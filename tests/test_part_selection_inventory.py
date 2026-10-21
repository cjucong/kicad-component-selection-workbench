import json
import shutil
import subprocess
import threading
from http.server import ThreadingHTTPServer
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from pathlib import Path

import openpyxl
import pytest
from unittest.mock import patch

from part_selection_inventory import (
    calculate_inventory_status,
    load_inventory,
    merge_inventory_records,
    parse_inventory_workbook,
    save_inventory,
    apply_inventory_matches,
)
from part_selection_server import PartSelectionApp, render_index


def _workbook(path: Path, rows, *, sheet_name="商品明细"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["商品编号", "厂家型号", "品牌", "封装", "订单编号", "同型号已购合计"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_import_uses_current_aggregate_and_deduplicates_equal_order_rows(scratch_project):
    source = scratch_project / "orders.xlsx"
    _workbook(source, [
        ["C100", "ABC-1", "Acme", "0603", "SO1", 12],
        ["C100", "ABC-1", "Acme", "0603", "SO2", 12],
    ])

    preview = parse_inventory_workbook(source)

    assert preview["items"]["C100"]["quantity"] == 12
    assert preview["items"]["C100"]["status"] == "ready"
    assert preview["items"]["C100"]["source_orders"] == ["SO1", "SO2"]
    assert preview["summary"]["ready"] == 1


def test_import_marks_conflicting_aggregate_for_review_instead_of_adding(scratch_project):
    source = scratch_project / "orders.xlsx"
    _workbook(source, [
        ["C100", "ABC-1", "Acme", "0603", "SO1", 12],
        ["C100", "ABC-1", "Acme", "0603", "SO2", 20],
    ])

    item = parse_inventory_workbook(source)["items"]["C100"]

    assert item["status"] == "unresolved"
    assert item["quantity"] is None
    assert item["source_orders"] == ["SO1", "SO2"]


def test_inventory_upsert_preserves_omitted_codes_and_repeat_import_does_not_add():
    existing = {"schema_version": 1, "items": {"C200": {"quantity": 4}}}
    preview = {"items": {"C100": {"quantity": 12, "status": "ready"}}}

    merged = merge_inventory_records(existing, preview)
    merged_again = merge_inventory_records(merged, preview)

    assert merged_again["items"]["C100"]["quantity"] == 12
    assert merged_again["items"]["C200"]["quantity"] == 4


def test_supplier_codes_are_preserved_pending_and_shared_model_total_is_not_added(scratch_project):
    source = scratch_project / 'supplier.xlsx'
    _workbook(source, [["D100", "MUX1", "Acme", "QFN", "SO1", 36],
                       ["D200", "MUX1", "Acme", "QFN", "SO2", 36]])
    preview = parse_inventory_workbook(source)
    assert preview['items'] == {}
    assert len(preview['unmatched']) == 1
    key, pending = next(iter(preview['unmatched'].items()))
    assert pending['quantity'] == 36
    assert pending['source_product_codes'] == ['D100', 'D200']
    persisted = merge_inventory_records({'items': {}}, preview)
    assert persisted['unmatched'][key]['quantity'] == 36
    matched = apply_inventory_matches(preview, {key: 'C700'})
    assert matched['items']['C700']['quantity'] == 36
    assert not matched['unmatched']
    assert not merge_inventory_records(persisted, matched).get('unmatched')


def test_manual_mapping_conflicting_with_direct_code_requires_review(scratch_project):
    source = scratch_project / 'supplier.xlsx'
    _workbook(source, [["C700", "MUX1", "Acme", "QFN", "SO1", 9],
                       ["D200", "MUX1", "Acme", "QFN", "SO2", 36]])
    preview = parse_inventory_workbook(source)
    key = next(iter(preview['unmatched']))
    matched = apply_inventory_matches(preview, {key: 'C700'})
    assert matched['items']['C700']['status'] == 'unresolved'
    assert matched['items']['C700']['quantity'] is None
    with pytest.raises(ValueError):
        apply_inventory_matches(preview, {key: 'D123'})


def test_supplier_source_cannot_be_bound_to_two_inventory_codes(inventory_app):
    app = inventory_app
    source = app.project_root / 'supplier.xlsx'
    _workbook(source, [["D100", "MUX1", "Acme", "QFN", "SO1", 36]])
    preview = app.preview_local_inventory({'path': str(source)})
    key = next(iter(preview['unmatched']))
    matched = app.preview_local_inventory({'path': str(source), 'matches': {key: 'C700'}})
    app.import_local_inventory({'preview_id': matched['preview_id'], 'confirmed': True})
    again = app.preview_local_inventory({'path': str(source)})
    assert not again['unmatched']
    assert again['items']['C700']['quantity'] == 36
    with pytest.raises(ValueError, match='绑定|归属'):
        app.preview_local_inventory({'path': str(source), 'matches': {key: 'C800'}})
    parsed = parse_inventory_workbook(source)
    other = apply_inventory_matches(parsed, {key: 'C800'})
    with pytest.raises(ValueError, match='绑定|归属'):
        merge_inventory_records(load_inventory(app.local_inventory_path), other)
    assert 'C800' not in load_inventory(app.local_inventory_path)['items']


def test_invalid_supplier_quantity_blocks_entire_model_group(scratch_project):
    source = scratch_project / 'supplier.xlsx'
    _workbook(source, [["D100", "MUX1", "Acme", "QFN", "SO1", 36],
                       ["D200", "MUX1", "Acme", "QFN", "SO2", None]])
    preview = parse_inventory_workbook(source)
    key, pending = next(iter(preview['unmatched'].items()))
    assert pending['status'] == 'unresolved'
    assert pending['quantity'] is None
    assert pending['source_product_codes'] == ['D100', 'D200']
    matched = apply_inventory_matches(preview, {key: 'C700'})
    assert matched['items']['C700']['status'] == 'unresolved'


def test_inventory_records_round_trip_atomically(scratch_project):
    path = scratch_project / "本地库存.json"
    inventory = {"schema_version": 1, "items": {"C100": {"quantity": 12}}}

    save_inventory(path, inventory)

    assert load_inventory(path) == inventory
    assert json.loads(path.read_text(encoding="utf-8")) == inventory
    assert list(scratch_project.iterdir()) == [path]


@pytest.mark.parametrize("sets", [0, -1, 1.5, "", 1000001])
def test_invalid_set_counts_are_rejected(sets):
    with pytest.raises(ValueError):
        calculate_inventory_status({"quantity": 5}, 2, sets)


def test_set_demand_and_shortage_are_calculated_from_group_quantity():
    status = calculate_inventory_status({"quantity": 5, "status": "ready"}, 2, 3)

    assert status == {
        "inventory": 5,
        "per_set_quantity": 2,
        "set_count": 3,
        "required": 6,
        "difference": -1,
        "status": "insufficient",
    }


def test_config_defaults_local_inventory_to_project_selection_data_directory(scratch_project):
    config = scratch_project / "BOM" / "选料工作台" / "工作台配置.json"
    config.parent.mkdir(parents=True)
    config.write_text(json.dumps({"project_root": "../..", "source_bom": "__schematic__"}), encoding="utf-8")
    app = PartSelectionApp(config)

    assert app.local_inventory_path == config.parent / "本地库存.json"


def test_import_preview_does_not_write_and_confirmation_upserts(scratch_project):
    config = scratch_project / "BOM" / "选料工作台" / "工作台配置.json"
    config.parent.mkdir(parents=True)
    config.write_text(json.dumps({"project_root": "../..", "source_bom": "__schematic__"}), encoding="utf-8")
    app = PartSelectionApp(config)
    source = scratch_project / "orders.xlsx"
    _workbook(source, [["C100", "ABC-1", "Acme", "0603", "SO1", 12]])

    preview = app.preview_local_inventory({"path": str(source)})
    assert not app.local_inventory_path.exists()
    with pytest.raises(ValueError, match="确认"):
        app.import_local_inventory({"path": str(source), "sheet": preview["sheet"]})
    imported = app.import_local_inventory({"preview_id": preview["preview_id"], "confirmed": True})
    preview = app.preview_local_inventory({"path": str(source)})
    repeated = app.import_local_inventory({"preview_id": preview["preview_id"], "confirmed": True})

    assert imported["summary"]["added"] == 1
    assert repeated["summary"]["repeated"] == 1
    assert load_inventory(app.local_inventory_path)["items"]["C100"]["quantity"] == 12


def test_rendered_workbench_adds_local_inventory_column_and_import_menu(scratch_project):
    config = scratch_project / "BOM" / "选料工作台" / "工作台配置.json"
    config.parent.mkdir(parents=True)
    config.write_text(json.dumps({"project_root": "../..", "source_bom": "__schematic__"}), encoding="utf-8")
    html = render_index(PartSelectionApp(config))

    assert '["local_inventory", "本地库存"]' in html
    assert 'data-col="local_inventory"' in html
    assert 'data-command="importLocalInventory"' in html
    assert '/api/local_inventory/preview' in html


@pytest.fixture
def inventory_app(scratch_project):
    config = scratch_project / "settings.json"
    config.write_text(json.dumps({"bom_dir": "custom-data", "schematic_dir": "."}))
    app = PartSelectionApp(config)
    yield app
    app.search_jobs.close()
    app.selection_jobs.close()


def test_preview_token_rejects_unknown_source_change_and_changed_inventory(inventory_app):
    app = inventory_app
    assert app.local_inventory_path == app.project_root / "custom-data/本地库存.json"
    source = app.project_root / "orders.xlsx"
    _workbook(source, [["C100", "ABC", "Acme", "0603", "SO1", 12]])
    preview = app.preview_local_inventory({"path": str(source)})
    with pytest.raises(ValueError):
        app.import_local_inventory({"confirmed": True, "preview_id": "unknown"})
    _workbook(source, [["C100", "ABC", "Acme", "0603", "SO1", 99]])
    with pytest.raises(ValueError, match="变|预览"):
        app.import_local_inventory({"confirmed": True, "preview_id": preview["preview_id"]})
    preview = app.preview_local_inventory({"path": str(source)})
    save_inventory(app.local_inventory_path, {"schema_version": 1, "items": {"C200": {"quantity": 5}}})
    with pytest.raises(ValueError, match="变|预览"):
        app.import_local_inventory({"confirmed": True, "preview_id": preview["preview_id"]})


def test_pending_inventory_can_be_explicitly_matched_in_preview(inventory_app):
    app = inventory_app
    source = app.project_root / 'supplier.xlsx'
    _workbook(source, [["D100", "MUX1", "Acme", "QFN", "SO1", 36],
                       ["D200", "MUX1", "Acme", "QFN", "SO2", 36]])
    preview = app.preview_local_inventory({'path': str(source)})
    key = next(iter(preview['unmatched']))
    app.import_local_inventory({'preview_id': preview['preview_id'], 'confirmed': True})
    assert load_inventory(app.local_inventory_path)['items'] == {}
    matched = app.preview_local_inventory({'path': str(source), 'matches': {key: 'C700'}})
    app.import_local_inventory({'preview_id': matched['preview_id'], 'confirmed': True})
    saved = load_inventory(app.local_inventory_path)
    assert saved['items']['C700']['quantity'] == 36
    assert not saved.get('unmatched')


def test_shared_codes_aggregate_demand_and_nc_does_not_consume_inventory(inventory_app):
    app = inventory_app
    save_inventory(app.local_inventory_path, {"schema_version": 1, "items": {
        "C100": {"quantity": 10, "status": "ready", "model": "SAMPLE"}}})
    rows = [{"quantity": 2, "effective_c_code": "C100"},
            {"quantity": 3, "effective_c_code": "C100"},
            {"quantity": 9, "effective_c_code": "C100", "value": "NC"},
            {"quantity": 1, "effective_c_code": "", "comment": "SAMPLE"},
            {"quantity": 1, "effective_c_code": "C100", "selection_codes_mixed": True}]
    app.enrich_local_inventory(rows)
    assert rows[0]["local_inventory"]["per_set_quantity"] == 5
    assert rows[0]["local_inventory"]["required"] == 15
    assert rows[0]["local_inventory"]["status"] == "insufficient"
    assert rows[2]["local_inventory"]["status"] == "nc"
    assert rows[3]["local_inventory"]["status"] == "unmatched"
    assert rows[3]["local_inventory"]["model_suggestions"] == ["C100"]
    assert rows[4]["local_inventory"]["status"] == "unresolved"


def test_workbook_aliases_invalid_rows_conflicts_and_missing_formula_cache(inventory_app):
    path = inventory_app.project_root / "fixture.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "说明"
    sheet = workbook.create_sheet("库存")
    sheet.append(["C 编码", "现存数量", "型号"])
    for row in [["C100", 0, "A"], ["wrong C200", 1, "B"], ["C300", -2, "C"],
                ["C400", "=1+1", "D"], ["C500", 4, "E"], ["C500", 5, "E"]]:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    preview = parse_inventory_workbook(path)
    assert preview["sheet"] == "库存"
    assert preview["items"]["C100"]["quantity"] == 0
    assert "C200" not in preview["items"]
    assert preview["items"]["C400"]["status"] == "unresolved"
    assert preview["items"]["C500"]["status"] == "unresolved"
    assert preview["summary"]["invalid"] == 3
    with pytest.raises(ValueError):
        parse_inventory_workbook(path, "说明")
    merged = merge_inventory_records({"items": {"C500": {"quantity": 99}}}, preview)
    assert merged["items"]["C500"]["status"] == "unresolved"


def test_inventory_read_error_is_reported_without_hiding_bom(inventory_app):
    inventory_app.local_inventory_path.parent.mkdir()
    inventory_app.local_inventory_path.write_text('invalid-json')
    rows = [{"quantity": 2, "effective_c_code": "C100"}]
    inventory_app.enrich_local_inventory(rows)
    assert rows[0]["local_inventory"]["status"] == "unresolved"
    assert rows[0]["local_inventory"]["issue"]


def test_inventory_http_preview_confirm_read_and_error(inventory_app):
    from part_selection_server import make_handler
    app = inventory_app
    source = app.project_root / 'fixture.xlsx'
    _workbook(source, [["C100", "ABC", "Acme", "0603", "SO1", 12]])
    server = ThreadingHTTPServer(('127.0.0.1', 0), make_handler(app))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f'http://127.0.0.1:{server.server_port}'
    def request(route, data=None):
        req = Request(base + route, data=json.dumps(data).encode() if data is not None else None,
                      headers={'Content-Type': 'application/json'})
        with urlopen(req, timeout=3) as response:
            return json.load(response)
    try:
        assert request('/api/local_inventory')['items'] == {}
        preview = request('/api/local_inventory/preview', {'path': str(source)})
        assert not app.local_inventory_path.exists()
        with pytest.raises(HTTPError) as error:
            request('/api/local_inventory/import', {'preview_id': preview['preview_id']})
        assert error.value.code == 400
        assert request('/api/local_inventory/import', {'preview_id': preview['preview_id'], 'confirmed': True})['ok']
        assert request('/api/local_inventory')['items']['C100']['quantity'] == 12
        with pytest.raises(HTTPError):
            request('/api/local_inventory/import', {'preview_id': preview['preview_id'], 'confirmed': True})
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def test_inventory_path_moves_with_registered_data_paths(inventory_app):
    from part_selection_settings import settings_snapshot
    app = inventory_app
    save_inventory(app.local_inventory_path, {'schema_version': 1, 'items': {'C100': {'quantity': 4}}})
    old = app.local_inventory_path
    app.save_path_settings({'revision': settings_snapshot(app.config)['revision'],
        'values': {'bom_dir': 'new-data', 'local_inventory': 'new-data/本地库存.json'}, 'copy_existing': True})
    assert app.local_inventory_path == app.project_root / 'new-data/本地库存.json'
    assert load_inventory(old) == load_inventory(app.local_inventory_path)


def test_atomic_replace_failure_leaves_previous_inventory(inventory_app):
    path = inventory_app.local_inventory_path
    original = {'schema_version': 1, 'items': {'C100': {'quantity': 4}}}
    save_inventory(path, original)
    with patch('part_selection_inventory.os.replace', side_effect=OSError('disk unavailable')):
        with pytest.raises(OSError):
            save_inventory(path, {'schema_version': 1, 'items': {}})
    assert load_inventory(path) == original
    assert not list(path.parent.glob('.inventory-*'))


@pytest.mark.parametrize('headers,values', [
    (['C编码', '已购合计', '同型号已购合计'], ['C100', 100, 12]),
    (['C编码', '同型号已购合计', '已购合计'], ['C100', 12, 100]),
])
def test_current_inventory_header_wins_over_legacy_quantity_column(scratch_project, headers, values):
    source = scratch_project / 'columns.xlsx'
    workbook = openpyxl.Workbook()
    workbook.active.append(headers)
    workbook.active.append(values)
    workbook.save(source)
    workbook.close()
    assert parse_inventory_workbook(source)['items']['C100']['quantity'] == 12


def test_implicit_inventory_path_change_requires_copy_and_conflict_checks(inventory_app):
    from part_selection_settings import settings_snapshot
    app = inventory_app
    save_inventory(app.local_inventory_path, {'schema_version': 1, 'items': {'C100': {'quantity': 4}}})
    old = app.local_inventory_path
    payload = {'revision': settings_snapshot(app.config)['revision'], 'values': {'bom_dir': 'moved'}}
    with pytest.raises(ValueError, match='复制'):
        app.save_path_settings(payload)
    target = app.project_root / 'moved/本地库存.json'
    save_inventory(target, {'schema_version': 1, 'items': {'C200': {'quantity': 9}}})
    with pytest.raises(ValueError, match='不同数据'):
        app.save_path_settings({**payload, 'copy_existing': True})
    target.unlink()
    app.save_path_settings({**payload, 'copy_existing': True})
    assert app.local_inventory_path == target
    assert load_inventory(old) == load_inventory(target)


@pytest.mark.skipif(not shutil.which('node'), reason='Node required for frontend inventory calculations')
def test_real_frontend_inventory_calculation_handles_shared_codes_and_unknown_zero():
    source = Path(__file__).resolve().parents[1] / 'tools/web/local-inventory.js'
    result = subprocess.run([shutil.which('node'), '-e', r'''
const assert=require('node:assert/strict');
const {localInventoryComparison:compare}=require(process.argv[1]);
assert.equal(compare({quantity:10,status:'ready'},2,5,3).difference,-5);
assert.equal(compare({quantity:10,status:'ready'},2,5,3).rowRequired,6);
assert.equal(compare({quantity:10,status:'ready'},2,5,3).shared,true);
assert.equal(compare({quantity:0,status:'ready'},2,2,3).status,'insufficient');
assert.equal(compare(null,2,2,3).quantity,null);
assert.equal(compare(null,2,2,3).status,'unmatched');
assert.equal(compare({quantity:10,status:'unresolved'},2,2,3).status,'unresolved');
assert.equal(compare({quantity:10},2,2,3,{mixed:true}).status,'unresolved');
assert.equal(compare({quantity:10},2,2,3,{nc:true}).status,'nc');
for(const count of [0,1.5,NaN,1000001]) assert.throws(()=>compare(null,1,1,count));
''', str(source)], capture_output=True, text=True, encoding='utf-8', timeout=10)
    assert result.returncode == 0, result.stderr


@pytest.mark.skipif(not shutil.which('node'), reason='Node required for inventory presentation')
def test_inventory_messages_use_available_supply_and_never_claim_unknown_stock_sufficient():
    source = Path(__file__).resolve().parents[1] / 'tools/web/local-inventory.js'
    result = subprocess.run([shutil.which('node'), '-e', r'''
const assert=require('node:assert/strict');
const {localInventoryComparison:compare,localInventoryPresentation:present}=require(process.argv[1]);
const sufficient=compare({quantity:20,status:'ready'},2,4,3);
assert.equal(present(sufficient,{},true).label,'现货充足');
const zero=compare({quantity:0,status:'ready'},2,4,3);
assert.equal(present(zero,{ok:true,stock:'12'},true).label,'现货充足');
assert.equal(present(zero,{ok:true,stock:'11'},true).label,'现货不足');
assert.equal(present(zero,{ok:false,stock:'999'},true).label,'商城现货待查询');
assert.equal(present(zero,{ok:true,stock:''},true).label,'商城现货待查询');
assert.equal(present(zero,{ok:true,stock:'26K'},true).label,'现货充足');
assert.equal(present(zero,{ok:true,stock:'0'},true).label,'现货不足');
assert.equal(present(zero,{ok:true,stock:'999',c_code:'C100'},true,'C200').label,'商城现货待查询');
assert.equal(present(zero,{ok:true,stock:'12'},true).local,'本地库存为零（现存 0）');
const absent=compare(null,2,4,3);
assert.equal(present(absent,{ok:true,stock:'12'},true).label,'现货充足');
assert.equal(present(absent,{},false).label,'待选料');
assert.equal(present(compare({quantity:10,status:'unresolved'},2,4,3),{ok:true,stock:'99'},true).label,'本地库存待核实');
assert.equal(present(compare(null,2,4,3,{nc:true}),{},false).label,'NC，不计需求');
assert.equal(present(compare({quantity:8},2,4,3),{ok:true,stock:'4'},true).label,'现货充足');
''', str(source)], capture_output=True, text=True, encoding='utf-8', timeout=10)
    assert result.returncode == 0, result.stderr
