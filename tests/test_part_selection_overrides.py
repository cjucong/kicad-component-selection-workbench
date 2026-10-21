import csv
import json
import re
from _scratch import temporary_project
import threading
import time
import unittest
from unittest import mock
from pathlib import Path

import sys

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TOOLS_DIR = PROJECT_ROOT / "tools"
sys.path.insert(0, str(TOOLS_DIR))

from part_selection_overrides import (  # noqa: E402
    choose_effective_c_code,
    load_config,
    load_overrides,
    normalize_alternate_parts,
    normalize_procurement_choice,
    save_override,
)
from part_selection_helper import PartInfo  # noqa: E402
from part_selection_helper import load_lcsc_bom, load_source_bom, read_refs_from_bom  # noqa: E402
from part_selection_helper import update_schematic_symbol_fields  # noqa: E402
import part_selection_server  # noqa: E402
from part_selection_runtime import runtime_directory
from part_selection_server import (  # noqa: E402
    PartSelectionApp,
    alternate_parts_from_schematic,
    alternate_parts_to_schematic_fields,
    group_parts_for_review,
    procurement_choice_from_schematic,
    procurement_choice_to_schematic_fields,
)


class PartSelectionOverridesTest(unittest.TestCase):
    def render_minimal_index(self):
        tmp = temporary_project()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        (root / "BOM").mkdir()
        config_path = root / "BOM" / "config.json"
        config_path.write_text(json.dumps({
            "source_bom": "BOM/source.csv",
            "schematic_dir": "DOC",
            "output_overrides": "BOM/overrides.json",
        }), encoding="utf-8")
        (root / "BOM" / "source.csv").write_text(
            "Comment,Description,Designator,Footprint,Value\nCapacitor,Capacitor,C1,0402,10pF\n",
            encoding="utf-8",
        )
        return part_selection_server.render_index(PartSelectionApp(config_path))

    def test_stock_parameter_summary_includes_brand(self):
        self.assertEqual(
            part_selection_server.stock_param_text(
                {
                    "param": "阻值: 10kΩ",
                    "param_detail": "品牌: VISHAY; 精度: ±1%",
                    "brand": "VISHAY",
                    "package": "0603",
                    "arrange": "圆盘",
                    "category": "贴片电阻",
                }
            ),
            "阻值: 10kΩ\n品牌: VISHAY; 精度: ±1%\n封装: 0603",
        )

    def test_sync_records_save_all_designators_once_and_preserve_selection(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            config = root / "config.json"
            config.write_text(json.dumps({"schematic_dir": "DOC"}), encoding="utf-8")
            app = PartSelectionApp(config)
            save_override(app.overrides_path, "R1", {"accepted_c_code": "C123", "note": "keep"})
            expected = {f"R{i}": {"value": "10k", "footprint": "DSMU:0402"} for i in range(1, 1234)}
            with mock.patch("part_selection_overrides.save_overrides", wraps=part_selection_server.save_overrides) as save:
                app.record_schematic_sync_overrides({"updated_refs": [], "changes": {}}, expected)
            self.assertEqual(save.call_count, 1)
            actual = load_overrides(app.overrides_path)
            self.assertEqual(len(actual), 1233)
            self.assertEqual(actual["R1"]["accepted_c_code"], "C123")
            self.assertEqual(actual["R1"]["note"], "keep")
            self.assertEqual(actual["R1233"]["schematic_field_snapshot"], expected["R1233"])

    def test_override_save_keeps_previous_file_when_atomic_replace_fails(self):
        import part_selection_overrides
        with temporary_project() as tmp:
            path = Path(tmp) / "BOM" / "overrides.json"
            part_selection_overrides.save_overrides(path, {"R1": {"note": "original"}})
            before = path.read_bytes()
            with mock.patch("os.replace", side_effect=OSError(22, "Invalid argument")):
                with self.assertRaises(OSError):
                    part_selection_overrides.save_overrides(path, {"R1": {"note": "new"}})
            self.assertEqual(path.read_bytes(), before)
            self.assertFalse(list((Path(tmp) / "Temp").glob("overrides-*")))
            self.assertFalse((Path(tmp) / "Temp").exists())
            self.assertFalse(list(path.parent.glob(".overrides-*")))

    def test_load_config_resolves_paths_against_config_file(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            config_path = root / "BOM" / "part_selection_config.json"
            config_path.parent.mkdir()
            config_path.write_text(
                json.dumps(
                    {
                        "schematic_dir": "DOC",
                        "source_bom": "BOM/source.csv",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            config = load_config(config_path)

            self.assertEqual(config["project_root"], root)
            self.assertEqual(config["schematic_dir"], root / "DOC")
            self.assertEqual(config["source_bom"], root / "BOM" / "source.csv")
            self.assertEqual(config["output_overrides"], root / "BOM" / "overrides.json")

    def test_save_override_normalizes_ref_and_preserves_other_entries(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "overrides.json"
            save_override(path, "r1", {"accepted_c_code": "C1", "decision": "AcceptManual", "note": "first"})
            save_override(path, "R2", {"accepted_c_code": "C2", "decision": "AcceptLCSC"})

            data = load_overrides(path)

            self.assertEqual(sorted(data), ["R1", "R2"])
            self.assertEqual(data["R1"]["accepted_c_code"], "C1")
            self.assertEqual(data["R1"]["note"], "first")
            self.assertIn("updated_at", data["R2"])

    def test_review_confirmation_records_server_time_and_keeps_history_when_status_changes(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "overrides.json"
            confirmed = save_override(path, "R1", {"review_status": "Confirmed"})["R1"]
            timestamp = confirmed["review_confirmed_at"]
            self.assertTrue(timestamp)
            pending = save_override(path, "R1", {"review_status": "Pending", "review_confirmed_at": "forged"})["R1"]
            self.assertEqual(pending["review_confirmed_at"], timestamp)

    def test_save_override_preserves_structured_field_overrides(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "overrides.json"
            save_override(
                path,
                "r1",
                {"field_overrides": {"comment": "Resistor", "value": "0R", "footprint": "0402"}},
            )

            data = load_overrides(path)

            self.assertEqual(data["R1"]["field_overrides"]["comment"], "Resistor")
            self.assertEqual(data["R1"]["field_overrides"]["value"], "0R")
            self.assertEqual(data["R1"]["field_overrides"]["footprint"], "0402")
            self.assertIn("field_updated_at", data["R1"])

    def test_save_override_preserves_schematic_field_snapshot(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "overrides.json"
            save_override(
                path,
                "R1",
                {
                    "schematic_field_snapshot": {
                        "comment": "Resistor",
                        "value": "10k",
                        "footprint": "0402",
                    }
                },
            )

            snapshot = load_overrides(path)["R1"]["schematic_field_snapshot"]

        self.assertEqual(
            snapshot,
            {"comment": "Resistor", "value": "10k", "footprint": "0402"},
        )

    def test_save_override_preserves_normalized_alternate_parts(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "overrides.json"
            save_override(
                path,
                "r1",
                {
                    "alternate_parts": [
                        {"c_code": "c 1001", "model": "M1", "package": "0402", "param": "10k"},
                        {"c_code": "C1001", "model": "duplicate"},
                        {"c_code": "bad", "model": "invalid"},
                    ]
                },
            )

            data = load_overrides(path)

            self.assertEqual(
                data["R1"]["alternate_parts"],
                [{"c_code": "C1001", "model": "M1", "package": "0402", "param": "10k"}],
            )
            self.assertIn("alternate_parts_updated_at", data["R1"])

    def test_alternate_schematic_fields_round_trip(self):
        parts = [
            {"c_code": "C1001", "model": "M1", "package": "0402", "param": "10k | 1%"},
            {"c_code": "C1002", "model": "M2", "package": "0603", "param": "2W"},
        ]

        fields = alternate_parts_to_schematic_fields(parts)
        restored = alternate_parts_from_schematic(fields["lcsc_alt"], fields["lcsc_alt_info"])

        self.assertEqual(fields["lcsc_alt"], "C1001,C1002")
        self.assertEqual([item["c_code"] for item in restored], ["C1001", "C1002"])
        self.assertEqual(restored[0]["param"], "10k / 1%")

    def test_procurement_choice_persists_without_c_code_and_round_trips_schematic_fields(self):
        choice = {
            "source": "lcsc_partner",
            "supplier": "waldom",
            "supplier_sku": "G9330",
            "manufacturer_part_number": "AD4630-24BBCZ",
            "package": "BGA-64",
            "batch": "26+",
            "order_url": "https://item.szlcsc.com/details/RzkzMzA=/AD4630-24BBCZ.html",
            "related_lcsc_code": "C5136981",
        }
        with temporary_project() as tmp:
            path = Path(tmp) / "overrides.json"
            save_override(path, "U8", {"procurement_choice": choice})
            saved = load_overrides(path)["U8"]["procurement_choice"]

        self.assertEqual(saved["supplier_sku"], "G9330")
        self.assertEqual(saved["order_url"], choice["order_url"])
        self.assertNotIn("c_code", saved)
        self.assertTrue(saved["id"].startswith("order_"))

        fields = procurement_choice_to_schematic_fields(saved)
        restored = procurement_choice_from_schematic(fields["procurement_url"], fields["procurement_info"])
        self.assertEqual(fields["procurement_url"], choice["order_url"])
        self.assertEqual(restored["manufacturer_part_number"], "AD4630-24BBCZ")
        self.assertEqual(restored["related_lcsc_code"], "C5136981")
        self.assertEqual(normalize_procurement_choice(restored)["id"], saved["id"])

    def test_choose_effective_c_code_prefers_manual_then_engineering_then_lcsc(self):
        part = {"accepted_c_code": "C_ENG", "lcsc_match_c_code": "C_LCSC"}

        self.assertEqual(
            choose_effective_c_code(part, {"decision": "AcceptManual", "manual_c_code": "C_MANUAL"}),
            "C_MANUAL",
        )
        self.assertEqual(choose_effective_c_code(part, {"accepted_c_code": "C_ACCEPTED"}), "C_ACCEPTED")
        self.assertEqual(choose_effective_c_code(part, {"decision": "AcceptEngineering"}), "C_ENG")
        self.assertEqual(choose_effective_c_code(part, {"decision": "AcceptLCSC"}), "C_LCSC")
        self.assertEqual(choose_effective_c_code(part, {"decision": "Reject"}), "")
        self.assertEqual(choose_effective_c_code(part, {}), "C_ENG")

    def test_review_rows_group_by_comment_footprint_value_and_keep_page_refs(self):
        parts = [
            PartInfo(ref="R1", page="PA", comment="Resistor", value="100k", footprint="0402", accepted_c_code="C1"),
            PartInfo(ref="R2", page="LDO", comment="Resistor", value="100k", footprint="0402", accepted_c_code="C1"),
            PartInfo(ref="R3", page="PA", comment="Resistor", value="10k", footprint="0402", accepted_c_code="C2"),
        ]

        rows = group_parts_for_review(parts, {})

        self.assertEqual(len(rows), 2)
        first = rows[0]
        self.assertEqual(first["refs"], "R1,R2")
        self.assertEqual(first["quantity"], 2)
        self.assertEqual(first["page_refs"], "LDO: R2; PA: R1")
        self.assertEqual(first["comment"], "Resistor")
        self.assertEqual(first["footprint"], "0402")
        self.assertEqual(first["value"], "100k")
        self.assertEqual(
            first["ref_links"],
            [
                {"ref": "R1", "page": "PA", "url": "kicad-sch-search://R1?page=PA"},
                {"ref": "R2", "page": "LDO", "url": "kicad-sch-search://R2?page=LDO"},
            ],
        )

    def test_review_rows_apply_field_overrides_before_grouping(self):
        parts = [
            PartInfo(ref="R1", page="PA", comment="Resistor", value="0R", footprint="0402", accepted_c_code="C1"),
            PartInfo(ref="R2", page="PA", comment="Y16070R50000F9R", value="0R", footprint="0402", accepted_c_code="C1"),
        ]

        rows = group_parts_for_review(
            parts,
            {"R2": {"field_overrides": {"comment": "Resistor", "value": "0R", "footprint": "0402"}}},
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["refs"], "R1,R2")
        self.assertEqual(rows[0]["quantity"], 2)
        self.assertEqual(rows[0]["comment"], "Resistor")
        self.assertEqual(rows[0]["field_changed_refs"], "R2")
        self.assertEqual(rows[0]["field_edit_status"], "NeedsSchematicUpdate")
        self.assertEqual(
            rows[0]["field_diffs"]["comment"],
            {"old": "Y16070R50000F9R", "new": "Resistor", "refs": "R2"},
        )
        part_selection_server.enrich_rows_with_schematic_sync_state(
            rows,
            Path("."),
            actual_by_ref={
                "R1": {"comment": "Resistor", "value": "0R", "footprint": "0402", "lcsc_c_code": "C1",
                       "lcsc_url": "https://so.szlcsc.com/global.html?k=C1"},
                "R2": {"comment": "Y16070R50000F9R", "value": "0R", "footprint": "0402", "lcsc_c_code": "C1",
                       "lcsc_url": "https://so.szlcsc.com/global.html?k=C1"},
            },
        )
        self.assertEqual(rows[0]["field_diffs_by_ref"], {
            "R2": {"comment": {"old": "Y16070R50000F9R", "new": "Resistor"}}
        })

    def test_rows_adopts_external_schematic_footprint_after_synced_snapshot(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,OLD_FP,10k\n",
                encoding="utf-8",
            )
            (doc_dir / "PA.kicad_sch").write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:R")\n'
                '    (property "Reference" "R1" (at 0 0 0))\n'
                '    (property "Value" "10k" (at 0 0 0))\n'
                '    (property "Footprint" "NEW_FP" (at 0 0 0))\n'
                '    (property "Comment" "Resistor" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            save_override(
                bom_dir / "overrides.json",
                "R1",
                {
                    "field_overrides": {"footprint": "OLD_FP"},
                    "schematic_field_snapshot": {
                        "comment": "Resistor",
                        "value": "10k",
                        "footprint": "OLD_FP",
                    },
                    "schematic_synced_at": "2026-08-18T10:00:00",
                },
            )

            row = PartSelectionApp(config_path).rows("", "BOM/source.csv", include_stock=False)[0]
            override = load_overrides(bom_dir / "overrides.json")["R1"]

        self.assertEqual(row["footprint"], "NEW_FP")
        self.assertFalse(row["field_changed_refs"])
        self.assertEqual(row["external_schematic_update_refs"], "R1")
        self.assertEqual(override["schematic_field_snapshot"]["footprint"], "NEW_FP")
        self.assertNotIn("footprint", override.get("field_overrides", {}))

    def test_rows_preserves_pending_web_footprint_edit(self):
        row, override = self._row_with_web_and_schematic_footprints("OLD_FP", "WEB_FP")

        self.assertEqual(row["footprint"], "WEB_FP")
        self.assertEqual(row["field_changed_refs"], "R1")
        self.assertNotIn("field_conflicts", row)
        self.assertEqual(override["field_overrides"]["footprint"], "WEB_FP")

    def test_rows_marks_conflict_when_web_and_schematic_both_changed(self):
        row, override = self._row_with_web_and_schematic_footprints("KICAD_FP", "WEB_FP")

        self.assertEqual(row["footprint"], "WEB_FP")
        self.assertEqual(row["field_changed_refs"], "R1")
        self.assertEqual(row["field_conflict_refs"], "R1")
        conflict = row["field_conflicts"]["R1"]["footprint"]
        self.assertEqual(conflict["baseline"], "OLD_FP")
        self.assertEqual(conflict["web"], "WEB_FP")
        self.assertEqual(conflict["schematic"], "KICAD_FP")
        self.assertEqual(override["field_overrides"]["footprint"], "WEB_FP")

    def test_legacy_timestamp_comparison_handles_mixed_date_formats(self):
        self.assertFalse(
            part_selection_server._timestamp_not_after(
                "2026-08-18T12:00:00", "2026/8/17 20:16:13"
            )
        )
        self.assertTrue(
            part_selection_server._timestamp_not_after(
                "2026/8/17 10:37:36", "2026-08-17T16:22:59"
            )
        )

    def _row_with_web_and_schematic_footprints(self, schematic_footprint, web_footprint):
        temp_dir = temporary_project()
        self.addCleanup(temp_dir.cleanup)
        root = Path(temp_dir.name)
        bom_dir = root / "BOM"
        doc_dir = root / "DOC"
        bom_dir.mkdir()
        doc_dir.mkdir()
        (bom_dir / "source.csv").write_text(
            "Comment,Description,Designator,Footprint,Value\n"
            "Resistor,Resistor,R1,OLD_FP,10k\n",
            encoding="utf-8",
        )
        (doc_dir / "PA.kicad_sch").write_text(
            '(kicad_sch\n'
            '  (symbol (lib_id "Device:R")\n'
            '    (property "Reference" "R1" (at 0 0 0))\n'
            '    (property "Value" "10k" (at 0 0 0))\n'
            f'    (property "Footprint" "{schematic_footprint}" (at 0 0 0))\n'
            '    (property "Comment" "Resistor" (at 0 0 0))\n'
            '  )\n'
            ')\n',
            encoding="utf-8",
        )
        config_path = bom_dir / "part_selection_config.json"
        config_path.write_text(
            json.dumps(
                {
                    "source_bom": "BOM/source.csv",
                    "schematic_dir": "DOC",
                    "output_overrides": "BOM/overrides.json",
                }
            ),
            encoding="utf-8",
        )
        save_override(
            bom_dir / "overrides.json",
            "R1",
            {
                "field_overrides": {"footprint": web_footprint},
                "schematic_field_snapshot": {
                    "comment": "Resistor",
                    "value": "10k",
                    "footprint": "OLD_FP",
                },
            },
        )
        row = PartSelectionApp(config_path).rows("", "BOM/source.csv", include_stock=False)[0]
        return row, load_overrides(bom_dir / "overrides.json")["R1"]

    def test_review_rows_merge_different_schematic_alternate_lists_without_losing_per_ref_data(self):
        parts = [
            PartInfo(
                ref="R1",
                page="PA",
                comment="Resistor",
                value="10k",
                footprint="0402",
                lcsc_alt="C1001",
                lcsc_alt_info="C1001 | M1 | 0402 | 10k",
            ),
            PartInfo(
                ref="R2",
                page="PA",
                comment="Resistor",
                value="10k",
                footprint="0402",
                lcsc_alt="C1002",
                lcsc_alt_info="C1002 | M2 | 0402 | 10k",
            ),
        ]

        rows = group_parts_for_review(parts, {})

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["refs"], "R1,R2")
        self.assertEqual(rows[0]["alternate_parts"][0]["c_code"], "C1001")
        self.assertTrue(rows[0]["selection_mixed"])
        self.assertEqual(rows[0]["per_ref_selection"]["R2"]["alternate_parts"][0]["c_code"], "C1002")

    def test_update_schematic_symbol_fields_writes_selected_refs_and_reports_diffs(self):
        with temporary_project() as tmp:
            doc_dir = Path(tmp) / "DOC"
            doc_dir.mkdir()
            schematic = doc_dir / "PA.kicad_sch"
            schematic.write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C118" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '  )\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C387" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '    (property "LCSC" "C4857915" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )

            result = update_schematic_symbol_fields(
                doc_dir,
                ["C118"],
                {
                    "value": "100pF, C0G",
                    "footprint": "0805",
                    "comment": "Capacitor",
                    "lcsc_c_code": "C4857915",
                    "lcsc_url": "https://so.szlcsc.com/global.html?k=C4857915",
                    "description": "CAP CER 100pF 50V C0G 0805",
                    "lcsc_alt": "C1001,C1002",
                    "lcsc_alt_info": "C1001 | ALT1 | 0805 | 100pF\nC1002 | ALT2 | 0805 | 100pF",
                    "procurement_url": "https://item.szlcsc.com/details/RzkzMzA=/AD4630.html",
                    "procurement_info": "lcsc_partner | waldom | G9330 | AD4630-24BBCZ | BGA-64 | 26+ | C5136981",
                },
            )

            text = schematic.read_text(encoding="utf-8")
            self.assertTrue(result["ok"])
            self.assertEqual(result["updated_refs"], ["C118"])
            self.assertEqual(
                result["changes"]["C118"]["value"],
                {"old": "100pF", "new": "100pF, C0G"},
            )
            self.assertEqual(
                result["changes"]["C118"]["footprint"],
                {"old": "0402", "new": "0805"},
            )
            self.assertEqual(
                result["changes"]["C118"]["lcsc_c_code"],
                {"old": "", "new": "C4857915"},
            )
            self.assertEqual(
                result["changes"]["C118"]["lcsc_url"],
                {"old": "", "new": "https://so.szlcsc.com/global.html?k=C4857915"},
            )
            self.assertEqual(
                result["changes"]["C118"]["description"],
                {"old": "", "new": "CAP CER 100pF 50V C0G 0805"},
            )
            self.assertIn('(property "Value" "100pF, C0G"', text)
            self.assertIn('(property "Footprint" "0805"', text)
            self.assertIn('(property "LCSC" "C4857915"', text)
            self.assertIn('(property "LCSC_URL" "https://so.szlcsc.com/global.html?k=C4857915"', text)
            self.assertIn('(property "Description" "CAP CER 100pF 50V C0G 0805"', text)
            self.assertIn('(property "LCSC_ALT" "C1001,C1002"', text)
            self.assertIn('(property "LCSC_ALT_INFO" "C1001 | ALT1 | 0805 | 100pF\\nC1002 | ALT2 | 0805 | 100pF"', text)
            self.assertIn('(property "PROCUREMENT_URL" "https://item.szlcsc.com/details/RzkzMzA=/AD4630.html"', text)
            self.assertIn('(property "PROCUREMENT_INFO" "lcsc_partner | waldom | G9330 | AD4630-24BBCZ | BGA-64 | 26+ | C5136981"', text)
            self.assertIn('(property "Reference" "C387"', text)
            self.assertIn('(property "Value" "100pF" (at 0 0 0))', text)

    def test_update_schematic_symbol_fields_escapes_multiline_description(self):
        with temporary_project() as tmp:
            doc_dir = Path(tmp) / "DOC"
            doc_dir.mkdir()
            schematic = doc_dir / "PA.kicad_sch"
            schematic.write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C118" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '    (property "Description" "Capacitor" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )

            update_schematic_symbol_fields(
                doc_dir,
                ["C118"],
                {"description": "CAP CER 100pF 50V C0G 0805\n容值: 100pF\nTape & Reel (TR)"},
            )

            text = schematic.read_text(encoding="utf-8")
            self.assertIn(
                '(property "Description" "CAP CER 100pF 50V C0G 0805\\n容值: 100pF\\nTape & Reel (TR)"',
                text,
            )
            self.assertNotIn('(property "Description" "CAP CER 100pF 50V C0G 0805\n', text)

    def test_app_sync_schematic_writes_file_override_and_sync_log(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C118,0402,100pF\n",
                encoding="utf-8",
            )
            schematic = doc_dir / "PA.kicad_sch"
            schematic.write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C118" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            app.prepare_kicad_for_external_schematic_write = lambda _refs: {"ok": True, "message": "stubbed"}
            app.notify_kicad_schematic_changed = lambda _result: {"ok": True, "message": "stubbed"}

            result = app.sync_schematic(
                {
                    "refs": "C118",
                    "field_overrides": {
                        "comment": "Capacitor",
                        "value": "100pF, C0G",
                        "footprint": "0805",
                        "lcsc_c_code": "C4857915",
                        "description": "CAP CER 100pF 50V C0G 0805",
                        "lcsc_alt": "C1001",
                        "lcsc_alt_info": "C1001 | ALT-CAP | 0805 | 100pF 50V",
                    },
                }
            )

            overrides = load_overrides(bom_dir / "overrides.json")
            self.assertTrue(result["ok"])
            self.assertEqual(overrides["C118"]["field_overrides"]["value"], "100pF, C0G")
            self.assertIn("schematic_synced_at", overrides["C118"])
            self.assertIn("value: 100pF -> 100pF, C0G", overrides["C118"]["schematic_sync_summary"])
            schematic_text = schematic.read_text(encoding="utf-8")
            self.assertIn('(property "Value" "100pF, C0G"', schematic_text)
            self.assertIn('(property "LCSC" "C4857915"', schematic_text)
            self.assertIn('(property "LCSC_URL" "https://so.szlcsc.com/global.html?k=C4857915"', schematic_text)
            self.assertIn('(property "Description" "CAP CER 100pF 50V C0G 0805"', schematic_text)
            self.assertIn('(property "LCSC_ALT" "C1001"', schematic_text)
            self.assertIn('(property "LCSC_ALT_INFO" "C1001 | ALT-CAP | 0805 | 100pF 50V"', schematic_text)
            sync_log = bom_dir / "原理图同步记录.csv"
            self.assertTrue(sync_log.exists())
            log_text = sync_log.read_text(encoding="utf-8-sig")
            self.assertIn("C118", log_text)
            self.assertIn("100pF, C0G", log_text)
            self.assertIn("C4857915", log_text)
            self.assertIn("https://so.szlcsc.com/global.html?k=C4857915", log_text)
            self.assertIn("lcsc_alt", log_text)
            self.assertIn("C1001 | ALT-CAP | 0805 | 100pF 50V", log_text)
            runtime_log = runtime_directory(root) / "原理图同步运行日志.jsonl"
            runtime_records = [json.loads(line) for line in runtime_log.read_text(encoding="utf-8").splitlines()]
            self.assertEqual(result["runtime_log"], str(runtime_log))
            self.assertTrue(result["sync_attempt_id"])
            self.assertTrue(all(record["attempt_id"] == result["sync_attempt_id"] for record in runtime_records))
            self.assertIn("request_started", [record["event"] for record in runtime_records])
            self.assertIn("request_complete", [record["event"] for record in runtime_records])

    def test_app_alternate_actions_persist_order_limit_and_candidate_stock(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,10k\n",
                encoding="utf-8",
            )
            (doc_dir / "PA.kicad_sch").write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:R")\n'
                '    (property "Reference" "R1" (at 0 0 0))\n'
                '    (property "Value" "10k" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Resistor" (at 0 0 0))\n'
                '    (property "LCSC_ALT" "C1001" (at 0 0 0))\n'
                '    (property "LCSC_ALT_INFO" "C1001 | M1 | 0402 | 10k" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                        "stock_cache": "BOM/stock.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)

            app.save_alternate_parts(
                {
                    "refs": "R1",
                    "action": "add",
                    "c_code": "C1002",
                    "primary_c_code": "C9000",
                    "part": {
                        "c_code": "C1002",
                        "model": "M2",
                        "package": "0603",
                        "param": "10k 1%",
                        "stock": "234",
                        "delivery": "现货",
                    },
                }
            )
            app.save_alternate_parts(
                {"refs": "R1", "action": "move", "c_code": "C1002", "direction": "up"}
            )
            for code in ("C1003", "C1004", "C1005"):
                app.save_alternate_parts(
                    {"refs": "R1", "action": "add", "c_code": code, "part": {"c_code": code}}
                )

            with self.assertRaisesRegex(ValueError, "at most 5"):
                app.save_alternate_parts(
                    {"refs": "R1", "action": "add", "c_code": "C1006", "part": {"c_code": "C1006"}}
                )
            with self.assertRaisesRegex(ValueError, "primary"):
                app.save_alternate_parts(
                    {
                        "refs": "R1",
                        "action": "add",
                        "c_code": "C9000",
                        "primary_c_code": "C9000",
                        "part": {"c_code": "C9000"},
                    }
                )

            overrides = load_overrides(bom_dir / "overrides.json")
            self.assertEqual(
                [item["c_code"] for item in overrides["R1"]["alternate_parts"]],
                ["C1002", "C1001", "C1003", "C1004", "C1005"],
            )
            stock_cache = json.loads((bom_dir / "stock.json").read_text(encoding="utf-8"))
            self.assertEqual(stock_cache["C1002"]["stock"], "234")

            app.save_override({"refs": "R1", "stock_c_code": "C1002"})
            updated = load_overrides(bom_dir / "overrides.json")
            self.assertNotIn("C1002", [item["c_code"] for item in updated["R1"]["alternate_parts"]])

    def test_app_rows_marks_alternate_conflict_and_adopts_schematic(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,10k\n",
                encoding="utf-8",
            )
            (bom_dir / "overrides.json").write_text(
                json.dumps(
                    {
                        "R1": {
                            "alternate_parts": [
                                {"c_code": "C2002", "model": "WEB", "package": "0402", "param": "10k"}
                            ],
                            "procurement_choice": {
                                "source": "lcsc_partner",
                                "supplier": "web",
                                "supplier_sku": "WEB1",
                                "manufacturer_part_number": "WEB-MPN",
                                "order_url": "https://item.szlcsc.com/details/web/order.html"
                            }
                        }
                    }
                ),
                encoding="utf-8",
            )
            (doc_dir / "PA.kicad_sch").write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:R")\n'
                '    (property "Reference" "R1" (at 0 0 0))\n'
                '    (property "Value" "10k" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Resistor" (at 0 0 0))\n'
                '    (property "LCSC_ALT" "C1001" (at 0 0 0))\n'
                '    (property "LCSC_ALT_INFO" "C1001 | SCH | 0402 | 10k" (at 0 0 0))\n'
                '    (property "PROCUREMENT_URL" "https://item.szlcsc.com/details/schematic/order.html" (at 0 0 0))\n'
                '    (property "PROCUREMENT_INFO" "lcsc_partner | schematic | SCH1 | SCH-MPN | 0402 | 24+ | " (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)

            row = app.rows("", "BOM/source.csv", include_stock=False)[0]
            self.assertTrue(row["alternate_conflict"])
            self.assertTrue(row["procurement_conflict"])
            self.assertEqual(row["field_diffs"]["lcsc_alt"]["old"], "C1001")
            self.assertEqual(row["field_diffs"]["lcsc_alt"]["new"], "C2002")
            self.assertEqual(row["field_diffs"]["procurement_url"]["old"], "https://item.szlcsc.com/details/schematic/order.html")
            cleared_verify = app.verify_schematic_field_values(
                {"R1": {"procurement_url": "", "procurement_info": ""}}
            )
            self.assertFalse(cleared_verify["ok"])
            self.assertIn("procurement_url", cleared_verify["mismatches"]["R1"])

            app.save_alternate_parts({"refs": "R1", "action": "adopt_schematic"})
            app.save_procurement_choice({"refs": "R1", "action": "adopt_schematic"})
            adopted = app.rows("", "BOM/source.csv", include_stock=False)[0]
            self.assertFalse(adopted["alternate_conflict"])
            self.assertFalse(adopted["procurement_conflict"])
            self.assertEqual(adopted["alternate_parts"][0]["c_code"], "C1001")
            self.assertEqual(adopted["procurement_choice"]["supplier_sku"], "SCH1")

    def test_app_sync_schematic_runtime_log_records_failed_attempt(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            with self.assertRaisesRegex(ValueError, "ref is required"):
                PartSelectionApp(config_path).sync_schematic({"field_overrides": {"value": "1k"}})

            runtime_log = runtime_directory(root) / "原理图同步运行日志.jsonl"
            records = [json.loads(line) for line in runtime_log.read_text(encoding="utf-8").splitlines()]
            self.assertEqual(
                [record["event"] for record in records],
                ["request_started", "transaction_lock_acquired", "request_failed"],
            )
            self.assertEqual(records[-1]["details"]["message"], "ref is required")

    def test_app_sync_schematic_serializes_concurrent_transactions(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            state_lock = threading.Lock()
            state = {"active": 0, "max_active": 0}

            def sync_stub(_payload, _log_event):
                with state_lock:
                    state["active"] += 1
                    state["max_active"] = max(state["max_active"], state["active"])
                time.sleep(0.04)
                with state_lock:
                    state["active"] -= 1
                return {"ok": True, "updated_refs": [], "missing_refs": [], "schematic_verify": {"ok": True}}

            app._sync_schematic_impl = sync_stub
            threads = [
                threading.Thread(
                    target=app.sync_schematic,
                    args=({"ref": f"R{index}", "field_overrides": {"value": "1k"}},),
                )
                for index in range(2)
            ]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join()

            self.assertEqual(state["max_active"], 1)

    def test_run_kicad_script_hides_powershell_window(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            completed = mock.Mock(returncode=0, stdout="", stderr="")

            with mock.patch.object(part_selection_server.subprocess, "run", return_value=completed) as run:
                result = app.run_kicad_script("kicad-sch-search://R1", ["-NotifyOnly"])

            self.assertTrue(result["ok"])
            self.assertEqual(
                run.call_args.kwargs.get("creationflags", 0) & part_selection_server.subprocess.CREATE_NO_WINDOW,
                part_selection_server.subprocess.CREATE_NO_WINDOW,
            )
            command = run.call_args.args[0]
            self.assertIn("-NonInteractive", command)
            self.assertIn("-WindowStyle", command)
            self.assertIn("Hidden", command)
            startupinfo = run.call_args.kwargs["startupinfo"]
            self.assertEqual(startupinfo.wShowWindow, part_selection_server.subprocess.SW_HIDE)
            self.assertTrue(startupinfo.dwFlags & part_selection_server.subprocess.STARTF_USESHOWWINDOW)

    def test_app_sync_schematic_stops_when_kicad_window_did_not_close(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C118,0402,100pF\n",
                encoding="utf-8",
            )
            schematic = doc_dir / "PA.kicad_sch"
            original_text = (
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C118" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '  )\n'
                ')\n'
            )
            schematic.write_text(original_text, encoding="utf-8")
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            app.prepare_kicad_for_external_schematic_write = lambda _refs: {
                "ok": False,
                "message": "KiCad schematic window did not close",
                "returncode": 1,
            }
            app.notify_kicad_schematic_changed = lambda _result: {"ok": True, "message": "stubbed"}

            with self.assertRaisesRegex(RuntimeError, "did not close"):
                app.sync_schematic(
                    {
                        "refs": "C118",
                        "field_overrides": {"value": "100pF, C0G"},
                    }
                )

            self.assertEqual(schematic.read_text(encoding="utf-8"), original_text)

    def test_app_sync_schematic_saves_then_reopens_and_searches_target_ref(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C118,0402,100pF\n",
                encoding="utf-8",
            )
            (doc_dir / "PA.kicad_sch").write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C118" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            calls = []
            app.run_kicad_script = lambda url, switches: calls.append((url, switches)) or {"ok": True}

            app.sync_schematic(
                {
                    "refs": "C118",
                    "field_overrides": {"comment": "Capacitor", "value": "100pF, C0G", "footprint": "0805"},
                }
            )

            self.assertEqual(calls[0][0], "kicad-sch-search://C118?page=PA")
            self.assertIn("-SaveActiveFrame", calls[0][1])
            self.assertIn("-SaveOnly", calls[0][1])
            self.assertIn("-CloseAfterSave", calls[0][1])
            self.assertNotIn("-NotifyOnly", calls[1][1])
            self.assertNotIn("FindDialog", calls[1][1])
            self.assertNotIn("-ReloadFromDisk", calls[1][1])

    def test_app_sync_schematic_retries_when_open_kicad_overwrites_disk_after_write(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C203,0805,2.2uF\n",
                encoding="utf-8",
            )
            schematic = doc_dir / "FGND_POWER.kicad_sch"
            old_text = (
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C203" (at 0 0 0))\n'
                '    (property "Value" "2.2uF" (at 0 0 0))\n'
                '    (property "Footprint" "0805" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '    (property "LCSC" "C49326714" (at 0 0 0))\n'
                '    (property "LCSC_URL" "https://so.szlcsc.com/global.html?k=C49326714" (at 0 0 0))\n'
                '    (property "Description" "old part" (at 0 0 0))\n'
                '  )\n'
                ')\n'
            )
            schematic.write_text(old_text, encoding="utf-8")
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            app.prepare_kicad_for_external_schematic_write = lambda _refs: {"ok": True, "message": "stubbed"}
            app.notify_kicad_schematic_changed = lambda _result: {"ok": True, "message": "stubbed"}
            original_update = part_selection_server.update_schematic_symbol_field_map
            calls = {"count": 0}

            def update_then_simulate_stale_kicad_save(doc_dir_arg, ref_fields_arg):
                calls["count"] += 1
                result = original_update(doc_dir_arg, ref_fields_arg)
                if calls["count"] == 1:
                    schematic.write_text(old_text, encoding="utf-8")
                return result

            part_selection_server.update_schematic_symbol_field_map = update_then_simulate_stale_kicad_save
            try:
                result = app.sync_schematic(
                    {
                        "refs": "C203",
                        "field_overrides": {
                            "comment": "Capacitor",
                            "value": "2.2uF",
                            "footprint": "0805",
                            "lcsc_c_code": "C48579622",
                            "lcsc_url": "https://so.szlcsc.com/global.html?k=C48579622",
                            "description": "new part",
                        },
                    }
                )
            finally:
                part_selection_server.update_schematic_symbol_field_map = original_update

            text = schematic.read_text(encoding="utf-8")
            self.assertTrue(result["ok"])
            self.assertTrue(result["schematic_verify"]["ok"])
            self.assertGreaterEqual(result["schematic_verify"]["attempts"], 2)
            self.assertGreaterEqual(calls["count"], 2)
            self.assertIn('(property "LCSC" "C48579622"', text)
            self.assertIn('(property "Description" "new part"', text)

    def test_app_sync_schematic_rewrites_when_kicad_notify_overwrites_disk(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C203,0805,2.2uF\n",
                encoding="utf-8",
            )
            schematic = doc_dir / "FGND_POWER.kicad_sch"
            old_text = (
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C203" (at 0 0 0))\n'
                '    (property "Value" "2.2uF" (at 0 0 0))\n'
                '    (property "Footprint" "0805" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '    (property "LCSC" "C49326714" (at 0 0 0))\n'
                '    (property "LCSC_URL" "https://so.szlcsc.com/global.html?k=C49326714" (at 0 0 0))\n'
                '    (property "Description" "old part" (at 0 0 0))\n'
                '  )\n'
                ')\n'
            )
            schematic.write_text(old_text, encoding="utf-8")
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            prepare_calls = {"count": 0}

            def prepare_stub(_refs):
                prepare_calls["count"] += 1
                return {"ok": True, "message": "stubbed"}

            app.prepare_kicad_for_external_schematic_write = prepare_stub
            notify_calls = {"count": 0}

            def notify_then_simulate_stale_kicad_save(_result):
                notify_calls["count"] += 1
                if notify_calls["count"] == 1:
                    schematic.write_text(old_text, encoding="utf-8")
                return {"ok": True, "message": "stubbed"}

            app.notify_kicad_schematic_changed = notify_then_simulate_stale_kicad_save

            result = app.sync_schematic(
                {
                    "refs": "C203",
                    "field_overrides": {
                        "comment": "Capacitor",
                        "value": "2.2uF",
                        "footprint": "0805",
                        "lcsc_c_code": "C48579622",
                        "lcsc_url": "https://so.szlcsc.com/global.html?k=C48579622",
                        "description": "new part",
                    },
                }
            )

            text = schematic.read_text(encoding="utf-8")
            self.assertTrue(result["ok"])
            self.assertTrue(result["schematic_verify"]["ok"])
            self.assertEqual(result["schematic_verify"]["phase"], "after_notify")
            self.assertEqual(prepare_calls["count"], 2)
            self.assertEqual(notify_calls["count"], 2)
            runtime_log = runtime_directory(root) / "原理图同步运行日志.jsonl"
            runtime_events = [
                json.loads(line)["event"] for line in runtime_log.read_text(encoding="utf-8").splitlines()
            ]
            self.assertIn("recovery_started", runtime_events)
            self.assertIn("recovery_prepare_complete", runtime_events)
            self.assertIn("recovery_notify_verified", runtime_events)
            self.assertIn('(property "LCSC" "C48579622"', text)
            self.assertIn('(property "Description" "new part"', text)

    def test_app_bulk_sync_bom_state_remaps_refs_and_writes_stock_description(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "IC,Old,U56,LCC16,DG509B\n"
                "Resistor,Old,R404,0402,10k\n",
                encoding="utf-8",
            )
            (bom_dir / "lcsc.csv").write_text(
                "型号,封装,参数/规格,工程/已接受C编码,立创匹配C编码,位号\n"
                "DG509B,LCC16,DG509B,C2673167,C2673167,U56\n"
                "RC0402,0402,10k,C25744,C25744,R404\n",
                encoding="utf-8",
            )
            (bom_dir / "stock.json").write_text(
                json.dumps(
                    {
                        "C2673167": {
                            "ok": True,
                            "c_code": "C2673167",
                            "param": "1 500Ω UQFN-16 Analog Switches RoHS",
                            "param_detail": "工作温度: -40℃~+125℃",
                            "arrange": "Tape & Reel (TR)",
                            "category": "Analog Switches",
                        },
                        "C25744": {
                            "ok": True,
                            "c_code": "C25744",
                            "param": "RES 10k 0402",
                        },
                    }
                ),
                encoding="utf-8",
            )
            schematic = doc_dir / "Out_Stage.kicad_sch"
            schematic.write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:U")\n'
                '    (property "Reference" "U28" (at 0 0 0))\n'
                '    (property "Value" "OLD" (at 0 0 0))\n'
                '    (property "Footprint" "OLD" (at 0 0 0))\n'
                '    (property "Comment" "OLD" (at 0 0 0))\n'
                '  )\n'
                '  (symbol (lib_id "Device:R")\n'
                '    (property "Reference" "R404" (at 0 0 0))\n'
                '    (property "Value" "OLD" (at 0 0 0))\n'
                '    (property "Footprint" "OLD" (at 0 0 0))\n'
                '    (property "Comment" "OLD" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                        "stock_cache": "BOM/stock.json",
                    }
                ),
                encoding="utf-8",
            )
            app = PartSelectionApp(config_path)
            app.prepare_kicad_for_external_schematic_write = lambda _refs: {"ok": True, "message": "stubbed"}
            app.notify_kicad_schematic_changed = lambda _result, **_kwargs: {"ok": True, "message": "stubbed"}

            result = app.bulk_sync_bom_state_to_schematic(
                "BOM/lcsc.csv",
                ref_map={"U56": "U28"},
            )

            text = schematic.read_text(encoding="utf-8")
            self.assertTrue(result["ok"])
            self.assertEqual(result["rows"], 2)
            self.assertEqual(result["target_refs"], 2)
            self.assertEqual(result["remapped_refs"], {"U56": "U28"})
            self.assertEqual(result["missing_refs"], [])
            self.assertIn('(property "Reference" "U28"', text)
            self.assertNotIn('"U56"', text)
            self.assertIn('(property "Value" "DG509B"', text)
            self.assertIn('(property "Footprint" "LCC16"', text)
            self.assertIn('(property "LCSC" "C2673167"', text)
            self.assertIn('(property "LCSC_URL" "https://so.szlcsc.com/global.html?k=C2673167"', text)
            self.assertIn("1 500Ω UQFN-16 Analog Switches RoHS\\n工作温度: -40℃~+125℃", text)
            self.assertIn('(property "Reference" "R404"', text)
            self.assertIn('(property "LCSC" "C25744"', text)
            records = [json.loads(line) for line in Path(result["runtime_log"]).read_text(encoding="utf-8").splitlines()]
            events = [record["event"] for record in records]
            self.assertIn("bulk_request_started", events)
            self.assertIn("bulk_targets_ready", events)
            self.assertIn("bulk_write_verified", events)
            self.assertEqual(events[-1], "bulk_request_complete")
            self.assertTrue(all(record["attempt_id"] == result["sync_attempt_id"] for record in records))

    def test_bulk_sync_does_not_write_conflicting_codes_or_their_cached_descriptions(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            (root / "BOM").mkdir()
            (root / "DOC").mkdir()
            (root / "BOM" / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\nCapacitor,Old,C1,C0402,100nF\n"
                "Capacitor,Old,C2,C0402,100nF\n", encoding="utf-8")
            (root / "BOM" / "lcsc.csv").write_text(
                "型号,封装,参数/规格,工程/已接受C编码,立创匹配C编码,位号\n"
                "PART1,0402,param A,C123,C123,C1\nPART2,0402,param B,C456,C456,C2\n", encoding="utf-8")
            schematic = root / "DOC" / "page.kicad_sch"
            schematic.write_text('(kicad_sch\n'
                ' (symbol (lib_id "Device:C") (property "Reference" "C1" (at 0 0 0)) '
                '(property "Value" "old" (at 0 0 0)) (property "Footprint" "old" (at 0 0 0)) '
                '(property "Comment" "old" (at 0 0 0)))\n'
                ' (symbol (lib_id "Device:C") (property "Reference" "C2" (at 0 0 0)) '
                '(property "Value" "old" (at 0 0 0)) (property "Footprint" "old" (at 0 0 0)) '
                '(property "Comment" "old" (at 0 0 0)))\n)\n', encoding="utf-8")
            (root / "BOM" / "config.json").write_text(json.dumps({
                "source_bom": "BOM/source.csv", "schematic_dir": "DOC",
                "output_overrides": "BOM/overrides.json", "stock_cache": "BOM/stock.json"
            }), encoding="utf-8")
            cache = {code: {"ok": True, "c_code": code, "param": desc}
                     for code, desc in (("C123", "param A"), ("C456", "param B"))}
            (root / "BOM" / "stock.json").write_text(json.dumps(cache), encoding="utf-8")
            app = PartSelectionApp(root / "BOM" / "config.json")
            app.prepare_kicad_for_external_schematic_write = lambda _refs: {"ok": True}
            app.notify_kicad_schematic_changed = lambda *_args, **_kwargs: {"ok": True}
            result = app.bulk_sync_bom_state_to_schematic("BOM/lcsc.csv", notify=False)
            text = schematic.read_text(encoding="utf-8")
            self.assertEqual(result["target_refs"], 2)
            self.assertNotIn('"LCSC" "C123"', text)
            self.assertNotIn('"LCSC" "C456"', text)
            self.assertNotIn('"Description" "param A"', text)
            self.assertNotIn('"Description" "param B"', text)
            self.assertEqual(text.count('"Value" "old"'), 2)

    def test_bulk_sync_logs_failure_before_writing(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            config = root / "config.json"
            config.write_text(json.dumps({"schematic_dir": "DOC"}), encoding="utf-8")
            app = PartSelectionApp(config)
            with unittest.mock.patch.object(app, "rows", side_effect=RuntimeError("read failed")):
                with self.assertRaisesRegex(RuntimeError, "read failed") as caught:
                    app.bulk_sync_bom_state_to_schematic("__schematic__")
            records = [json.loads(line) for line in Path(caught.exception.runtime_log).read_text(encoding="utf-8").splitlines()]
            self.assertEqual(records[-1]["event"], "bulk_request_failed")
            self.assertEqual(records[-1]["details"]["message"], "read failed")
            self.assertEqual(records[-1]["attempt_id"], caught.exception.sync_attempt_id)

    def test_read_refs_from_xlsx_bom_designator_column(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "purchase.xlsx"
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.append(["封装", "参数/规格", "工程/已接受C编码", "立创匹配C编码", "位号"])
            ws.append(["0402", "100k", "C25741", "", "R1169,R1170"])
            ws.append(["0603", "10k", "", "C25804", "R1 R2"])
            wb.save(path)

            self.assertEqual(read_refs_from_bom(path), ["R1169", "R1170", "R1", "R2"])

    def test_xlsx_bom_rows_map_to_source_and_lcsc_data(self):
        with temporary_project() as tmp:
            path = Path(tmp) / "purchase.xlsx"
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.append(["封装", "参数/规格", "工程/已接受C编码", "立创匹配C编码", "位号", "页面位号", "单套用量"])
            ws.append(["0402", "100k", "C25741", "C25741", "R1169,R1170", "PA: R1169; LDO: R1170", 2])
            wb.save(path)

            source = load_source_bom(path)
            lcsc = load_lcsc_bom(path)

            self.assertEqual(source["R1169"]["Value"], "100k")
            self.assertEqual(source["R1170"]["Footprint"], "0402")
            self.assertEqual(lcsc["R1169"]["AcceptedCCode"], "C25741")
            self.assertEqual(lcsc["R1169"]["LCSCMatchCCode"], "C25741")
            self.assertEqual(lcsc["R1170"]["PageRefs"], "PA: R1169; LDO: R1170")
            self.assertEqual(lcsc["R1170"]["Quantity"], "2")

    def test_app_rows_loads_selected_bom_refs_when_filter_is_empty(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            bom_path = bom_dir / "selected.csv"
            bom_path.write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,\"R1,R2\",0402,100k\n"
                "Capacitor,Capacitor,C1,0402,100nF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/selected.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            app = PartSelectionApp(config_path)
            rows = app.rows("", "BOM/selected.csv")

            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["refs"], "C1")
            self.assertEqual(rows[1]["refs"], "R1,R2")
            self.assertEqual(rows[1]["quantity"], 2)

    def test_app_rows_sort_by_component_category_value_and_footprint(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            bom_path = bom_dir / "selected.csv"
            bom_path.write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "IC,IC,U1,QFN,ADC\n"
                "Resistor,Resistor,R10,0603,10k\n"
                "Capacitor,Capacitor,C2,0805,100nF\n"
                "Inductor,Inductor,L1,SMD-2F,33uH\n"
                "Capacitor,Capacitor,C1,0402,100pF\n"
                "Resistor,Resistor,R1,0402,1k\n"
                "Capacitor,Capacitor,C3,0402,100nF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/selected.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            rows = PartSelectionApp(config_path).rows("", "BOM/selected.csv", include_stock=False)

            self.assertEqual([row["refs"] for row in rows], ["C1", "C3", "C2", "R1", "R10", "L1", "U1"])

    def test_app_rows_can_use_schematic_as_direct_source(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            (doc_dir / "PA.kicad_sch").write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C118" (at 0 0 0))\n'
                '    (property "Value" "100pF" (at 0 0 0))\n'
                '    (property "Footprint" "0402" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '    (property "LCSC" "C4857915" (at 0 0 0))\n'
                '  )\n'
                '  (symbol (lib_id "Device:R")\n'
                '    (property "Reference" "R5" (at 0 0 0))\n'
                '    (property "Value" "10k" (at 0 0 0))\n'
                '    (property "Footprint" "0603" (at 0 0 0))\n'
                '    (property "Comment" "Resistor" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            app = PartSelectionApp(config_path)
            rows = app.rows("", "__schematic__", include_stock=False)

            self.assertEqual([row["refs"] for row in rows], ["C118", "R5"])
            self.assertEqual(rows[0]["value"], "100pF")
            self.assertEqual(rows[0]["accepted_c_code"], "C4857915")
            self.assertEqual(rows[0]["source_mode"], "schematic")

            boms = app.available_boms()
            self.assertEqual(boms[0]["path"], "__schematic__")

    def test_app_rows_reuses_cached_stock_for_last_queried_replacement_code(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0805,100nF\n",
                encoding="utf-8",
            )
            (bom_dir / "overrides.json").write_text(
                json.dumps({"C1": {"stock_c_code": "C5182461", "decision": "", "status": ""}}),
                encoding="utf-8",
            )
            (bom_dir / "stock.json").write_text(
                json.dumps(
                    {
                        "C5182461": {
                            "ok": True,
                            "status": "OK",
                            "c_code": "C5182461",
                            "stock": "4030",
                            "param": "CAP CER 100nF 250V X7R 0805",
                            "param_detail": "容值: 100nF; 温度系数: X7R; 精度: ±10%; 额定电压: 250V",
                        }
                    }
                ),
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                        "stock_cache": "BOM/stock.json",
                    }
                ),
                encoding="utf-8",
            )

            app = PartSelectionApp(config_path)
            rows = app.rows("", "BOM/source.csv", include_stock=False)

            self.assertEqual(rows[0]["stock_c_code"], "C5182461")
            self.assertEqual(rows[0]["per_ref_selection"]["C1"]["effective_c_code"], "C5182461")
            self.assertEqual(rows[0]["stock"]["c_code"], "C5182461")
            self.assertEqual(rows[0]["stock"]["param_detail"], "容值: 100nF; 温度系数: X7R; 精度: ±10%; 额定电压: 250V")

    def test_export_pump_table_filters_stock_issue_and_excludes_non_orderable_nc(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value,AcceptedCCode\n"
                "Capacitor,Capacitor,C1,0603,10uF,C1001\n"
                "Capacitor,Capacitor,C2,0603,1uF,C1002\n"
                "Capacitor,Capacitor,C3,0603,NC,C1003\n"
                "Resistor,Resistor,R1,0402,10k,\n",
                encoding="utf-8",
            )
            (bom_dir / "stock.json").write_text(
                json.dumps(
                    {
                        "C1001": {"ok": True, "c_code": "C1001", "stock": "15", "delivery": "现货", "param": "low stock cap"},
                        "C1002": {"ok": True, "c_code": "C1002", "stock": "200", "delivery": "现货", "param": "ok cap"},
                        "C1003": {"ok": True, "c_code": "C1003", "stock": "1", "delivery": "现货", "param": "nc cap"},
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                        "stock_cache": "BOM/stock.json",
                    }
                ),
                encoding="utf-8",
            )
            for ref, code in {"C1": "C1001", "C2": "C1002", "C3": "C1003"}.items():
                save_override(bom_dir / "overrides.json", ref, {"accepted_c_code": code, "stock_c_code": code})

            result = PartSelectionApp(config_path).export_pump_table({"bom": "BOM/source.csv", "export_filter": "low_stock"})

            self.assertTrue(result["ok"])
            self.assertEqual(result["count"], 1)
            with Path(result["path"]).open(encoding="utf-8-sig", newline="") as file:
                rows = list(csv.DictReader(file))
            self.assertEqual(rows[0]["Designator"], "C1")
            self.assertEqual(rows[0]["LCSC"], "C1001")
            self.assertNotIn("C3", [row["Designator"] for row in rows])
            self.assertNotIn("R1", [row["Designator"] for row in rows])

    def test_nc_detection_prefers_value_field_tokens(self):
        self.assertTrue(part_selection_server.row_is_nc({"comment": "Capacitor", "value": "NC"}))
        self.assertTrue(part_selection_server.row_is_nc({"comment": "Capacitor", "value": "10k, NC"}))
        self.assertTrue(part_selection_server.row_is_nc({"comment": "Resistor", "value": "DNP"}))
        self.assertFalse(part_selection_server.row_is_nc({"comment": "NC7SZ logic gate", "value": "SN74LVC1G17"}))
        self.assertFalse(part_selection_server.row_is_nc({"comment": "Capacitor", "value": "100nF"}))

    def test_nc_rows_are_not_stock_issue_filter_candidates(self):
        row = {
            "comment": "Capacitor",
            "value": "NC",
            "accepted_c_code": "C1003",
            "stock": {"ok": True, "stock": "1", "c_code": "C1003"},
        }

        keys = part_selection_server.row_filter_keys(row)

        self.assertIn("nc_device", keys)
        self.assertNotIn("stock_issue", keys)

    def test_row_filter_keys_include_inverse_filter_states(self):
        row = {
            "comment": "Resistor",
            "value": "10k",
            "review_status": "Confirmed",
            "stock": {"ok": True, "stock": "500", "c_code": "C1001"},
        }

        keys = part_selection_server.row_filter_keys(row)

        self.assertTrue(
            {
                "sync_done",
                "review_confirmed",
                "code_unchanged",
                "stock_normal",
                "ref_complete",
                "non_nc",
            }.issubset(keys)
        )

    def test_app_rows_marks_replacement_c_code_as_needing_schematic_sync(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            doc_dir = root / "DOC"
            bom_dir.mkdir()
            doc_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0603,10uF\n",
                encoding="utf-8",
            )
            (doc_dir / "PA.kicad_sch").write_text(
                '(kicad_sch\n'
                '  (symbol (lib_id "Device:C")\n'
                '    (property "Reference" "C1" (at 0 0 0))\n'
                '    (property "Value" "10uF" (at 0 0 0))\n'
                '    (property "Footprint" "0603" (at 0 0 0))\n'
                '    (property "Comment" "Capacitor" (at 0 0 0))\n'
                '    (property "LCSC" "C344022" (at 0 0 0))\n'
                '    (property "LCSC_URL" "https://so.szlcsc.com/global.html?k=C344022" (at 0 0 0))\n'
                '    (property "Description" "old part" (at 0 0 0))\n'
                '  )\n'
                ')\n',
                encoding="utf-8",
            )
            (bom_dir / "lcsc_stock_cache.json").write_text(
                json.dumps(
                    {
                        "C22392391": {
                            "ok": True,
                            "c_code": "C22392391",
                            "param": "CAP CER 10uF 25V X5R 0603",
                            "param_detail": "容值: 10uF; 温度系数: X5R; 额定电压: 25V",
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                        "stock_cache": "BOM/lcsc_stock_cache.json",
                    }
                ),
                encoding="utf-8",
            )
            save_override(
                bom_dir / "overrides.json",
                "C1",
                {
                    "decision": "AcceptManual",
                    "accepted_c_code": "C22392391",
                    "manual_c_code": "C22392391",
                    "stock_c_code": "C22392391",
                    "status": "ChangedToNewCode",
                },
            )

            rows = PartSelectionApp(config_path).rows("", "BOM/source.csv", include_stock=False)

            self.assertEqual(rows[0]["field_changed_refs"], "C1")
            self.assertEqual(rows[0]["field_edit_status"], "NeedsSchematicUpdate")
            self.assertIn("lcsc_c_code", rows[0]["field_diffs"])
            self.assertEqual(rows[0]["field_diffs"]["lcsc_c_code"]["old"], "C344022")
            self.assertEqual(rows[0]["field_diffs"]["lcsc_c_code"]["new"], "C22392391")

    def test_available_boms_excludes_generated_review_outputs(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            (bom_dir / "PXIE_4137_LCSC_BOM_20260813_101540.csv").write_text(
                "型号,封装,参数/规格,工程/已接受C编码,立创匹配C编码,位号\n"
                "RC0402,0402,100k,C25741,C25741,R1\n",
                encoding="utf-8",
            )
            (bom_dir / "part_selection_change_log_20260813_015410.csv").write_text(
                "Designator,OldCCode,NewCCode\n"
                "R1,C1,C2\n",
                encoding="utf-8",
            )
            (bom_dir / "part_selection_schematic_sync_log.csv").write_text(
                "SyncedAt,Designator,Page,Field,OldValue,NewValue,SchematicPath\n"
                "2026-08-13T10:00:00,R1,PA,value,10k,100k,DOC/PA.kicad_sch\n",
                encoding="utf-8",
            )

            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.append(["位号", "Comment", "Value"])
            ws.append(["R2", "Resistor", "10k"])
            wb.save(bom_dir / "part_selection_review_20260812_195951.xlsx")

            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            app = PartSelectionApp(config_path)
            paths = [item["path"] for item in app.available_boms()]

            self.assertEqual(
                paths,
                [
                    "__schematic__",
                    "BOM/PXIE_4137_LCSC_BOM_20260813_101540.csv",
                    "BOM/source.csv",
                ],
            )

    def test_render_index_removes_substitute_toggle_and_marks_candidate_columns(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertNotIn("显示参数替代", html)
            self.assertNotIn("includeSubstitutes", html)
            self.assertIn('data-candidate-col="action"', html)
            self.assertIn('data-candidate-col="param"', html)
            self.assertIn("选用主料", html)
            self.assertIn("加入备选", html)
            self.assertIn('data-col="alternates"', html)
            self.assertIn('["alternates", "备选料"]', html)
            self.assertIn("/api/alternate_parts", html)
            self.assertIn('data-action="move-alternate"', html)
            self.assertIn('data-action="adopt-schematic-alternates"', html)
            self.assertIn("const UI_SCHEMA_VERSION = 7", html)
            self.assertIn("badV4Default", html)
            self.assertIn("{note: true}", html)
            self.assertIn("vertical-align: middle", html)
            self.assertIn("min-height: 40px", html)
            self.assertIn("let defaultColumnWidthFrame = 0", html)
            self.assertIn("const latestSettings = loadUiSettings()", html)
            self.assertIn("window.addEventListener(\"storage\"", html)
            self.assertIn("candidate-section-row", html)
            self.assertIn("candidate-partner-row", html)
            self.assertIn("renderCandidatePartnerRows", html)
            self.assertIn('data-candidate-col="channel_stock"', html)
            self.assertIn("渠道 / 库存交期", html)
            self.assertNotIn('data-candidate-col="channel"', html)
            self.assertNotIn('data-candidate-col="stock"', html)
            self.assertIn("function renderCandidateParam", html)
            self.assertIn("function renderCandidateOrderLines", html)
            self.assertIn("function renderCandidatePrice", html)
            self.assertIn("candidate-param-key", html)
            self.assertIn('id="minimizeCandidate"', html)
            self.assertIn('id="maximizeCandidate"', html)
            self.assertIn('id="closeCandidate"', html)
            self.assertIn('id="candidateMinimizedBar"', html)
            self.assertEqual(html.count('class="modal-resize-handle"'), 8)
            self.assertIn("function beginCandidateMove", html)
            self.assertIn('event.target.closest(".modal-controls, .candidate-title")', html)
            self.assertIn('.addEventListener("pointerdown", beginCandidateMove)', html)
            self.assertIn("function beginCandidateResize", html)
            self.assertIn("targetKey === candidateTargetKey", html)
            self.assertIn("candidateWindow.classList.toggle(\"fullscreen\")", html)
            self.assertIn("function renderStockParam", html)
            self.assertIn("paramCell.innerHTML = renderStockParam(material)", html)
            self.assertIn("cell?.dataset.stockParamText", html)
            self.assertIn("sync-summary-detail", html)
            self.assertIn("let candidateSearchToken = 0", html)
            self.assertNotIn("event.target === candidateModal && candidateBackdropPointerDown", html)
            self.assertIn("合作库存", html)
            self.assertNotIn("${renderCandidateChannelSummary(item)}", html)
            self.assertIn("function candidateKeywordFromFields", html)
            self.assertIn("function isExplicitPassiveComment", html)
            self.assertIn("uniqueSearchTerms([comment, fields.value, fields.footprint])", html)
            self.assertIn('return String(fields.comment || fields.value || "").trim();', html)
            self.assertIn('const exactCCode = /^C\\d+$/.test(compactCode)', html)
            self.assertIn('include_substitutes=${exactCCode ? "0" : "1"}', html)
            self.assertIn("选用订货料", html)
            self.assertIn("立创订货", html)
            self.assertIn('/api/procurement_choice', html)
            self.assertIn('data-action="select-procurement"', html)
            self.assertIn('data-action="clear-procurement"', html)
            self.assertIn('data-action="adopt-schematic-procurement"', html)
            self.assertIn("PROCUREMENT_URL", html)
            self.assertNotIn('data-action="choose-candidate" data-code="${esc(code)}" title="采用该合作库存', html)

    def test_render_index_has_separate_missing_and_refresh_stock_actions(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn('id="queryMissingStock"', html)
            self.assertIn("查询未查库存", html)
            self.assertIn("刷新全部库存", html)
            self.assertIn('startStockQuery("missing"', html)
            self.assertIn('startStockQuery("all"', html)
            self.assertIn('id="stockProgress"', html)

    def test_render_index_shows_schematic_sync_progress_and_uses_bulk_api(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn('id="syncProgress"', html)
            self.assertIn('id="syncProgressFill"', html)
            self.assertIn('id="syncProgressText"', html)
            self.assertIn("/api/bulk_sync_schematic", html)
            self.assertIn("setSyncProgress", html)

    def test_render_index_marks_unsynced_schematic_fields_not_changed_c_code_rows(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn("action-sync-needed", html)
            self.assertIn("row-needs-sync", html)
            self.assertIn("schematicSyncMarker", html)
            self.assertIn("rowVisualClass(row, selectedStatus, syncNeeded, stock, reviewStatus)", html)

    def test_render_index_marks_c_code_changes_dirty_and_restores_saved_row(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0603,10uF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn("function markRowNeedsSchematicSync", html)
            self.assertIn("markRowNeedsSchematicSync(tr", html)
            self.assertIn("function rowAnchor", html)
            self.assertIn("restoreRowAnchor", html)
            self.assertIn("await loadRows(anchor)", html)

    def test_render_index_adds_row_number_and_part_confirmation_columns(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0603,10uF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn('["row_no", "行号"]', html)
            self.assertIn('["review", "确认料"]', html)
            self.assertIn('data-col="row_no"', html)
            self.assertNotIn('data-col="brand"', html)
            self.assertIn('const brand = String(stock.brand || "").trim();', html)
            self.assertIn('pieces.push(`封装: ${packageName}`)', html)
            self.assertNotIn('pieces.push(stock.arrange)', html)
            self.assertNotIn('esc(item.category || "")', html)
            self.assertIn('data-field="review_status"', html)
            self.assertIn("review-ok", html)
            self.assertIn("review-warn", html)
            self.assertIn("review-bad", html)
        self.assertIn("data-review-confirmed-at", html)
        self.assertIn("确认于", html)
        self.assertIn("上次确认于", html)

    def test_render_index_places_row_sync_in_review_and_model_search_in_stock(self):
        html = self.render_minimal_index()
        match = re.search(r'function renderRows\(rows\).*?(?=function updateRowsForStock)', html, re.S)
        self.assertIsNotNone(match)
        row_template = match.group()
        self.assertNotIn('data-col="search"', html)
        self.assertIn('data-action="open-candidates"', html)
        self.assertIn('data-action="sync-schematic"', row_template)
        self.assertIn('data-col="review">', row_template)
        review_cell = row_template.split('<td data-col="review">', 1)[1].split('</td>', 1)[0]
        self.assertIn('${syncButton}', review_cell)
        self.assertIn('data-action="undo-workbench"', review_cell)
        self.assertIn('${syncMarker}', review_cell)
        self.assertIn('data-action="sync-schematic"', html)
        self.assertIn('<td data-col="stock">', row_template)
        self.assertIn('<div data-stock-cell>', row_template)
        self.assertIn('function renderStockCell(stock)', html)
        self.assertIn('data-action="open-candidates">更换型号</button>', html)
        self.assertNotIn('["search", "搜索"]', html)
        self.assertNotIn('search: [170, 240]', html)
        self.assertIn('const UI_SCHEMA_VERSION = 7', html)
        self.assertIn('delete (settings.hiddenColumns || {}).search', html)
        self.assertIn('delete settings.columnWidths.search', html)
        self.assertIn('SYNC_FIELD_LABELS[field] || field', html)
        self.assertIn('values.old || "(空)"', html)
        self.assertIn('values.new || "(空)"', html)

    def test_inline_row_sync_uses_selected_bulk_sync_with_exact_row_refs(self):
        html = self.render_minimal_index()
        function = re.search(r'async function syncSchematicWithProgress\(tr\).*?(?=async function syncAllSchematicWithProgress)', html, re.S)
        self.assertIsNotNone(function)
        source = function.group()
        self.assertIn('postSchematicSync("/api/bulk_sync_schematic"', source)
        self.assertIn('const selectedRefs = syncRowRefs(tr)', source)
        self.assertIn('selected_refs: selectedRefs', source)
        self.assertIn('bom: loadedRowsBom', source)
        self.assertNotIn('postSchematicSync("/api/sync_schematic"', source)

    def test_render_index_has_filter_views_and_fixed_row_color_priority(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0603,10uF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            for filter_name in ("all", "needs_sync", "pending_review", "changed_code", "stock_issue", "missing_ref", "nc_device"):
                self.assertIn(f'data-filter="{filter_name}"', html)
            self.assertIn("function rowVisualClass(row, selectedStatus, syncNeeded, stock, reviewStatus)", html)
            self.assertIn('return "row-needs-sync"', html)
            self.assertIn('return "row-muted"', html)
            self.assertIn('return "row-warning"', html)
            self.assertIn('return "row-ok"', html)
            self.assertIn("function applyActiveFilter()", html)
            self.assertIn("data-filter-key", html)
            self.assertIn("const FILTER_CYCLES = {", html)
            self.assertIn("const activeFilters = new Map()", html)
            self.assertIn("keys.includes(filter)", html)
            for inverse_filter in (
                "sync_done",
                "review_confirmed",
                "code_unchanged",
                "stock_normal",
                "ref_complete",
                "non_nc",
            ):
                self.assertIn(inverse_filter, html)
            self.assertIn("Array.from(activeFilters.values())", html)
            self.assertIn("function updateFilterButtons()", html)

    def test_render_index_has_pump_table_export_controls(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0603,10uF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn('id="pumpExportMode"', html)
            self.assertIn('id="exportPumpTable"', html)
            self.assertIn("function organizeToolbar()", html)
            self.assertIn('class="toolbar-section source-section"', html)
            self.assertIn('class="toolbar-section action-section"', html)
            self.assertIn('class="toolbar-section export-section"', html)
            self.assertIn("filter-hint", html)
            self.assertIn("body.toolbar-compact .toolbar", html)
            self.assertIn("function updateStickyChrome()", html)
            self.assertIn("--sticky-toolbar-height", html)
            self.assertIn("main { padding: 0 0 24px; }", html)
            self.assertNotIn("main { padding: 0 0 24px; overflow-x: auto; }", html)
            self.assertIn("function observeStickyChrome()", html)
            self.assertIn("new ResizeObserver(scheduleStickyChromeUpdate)", html)
            self.assertIn("translateX(${window.scrollX}px)", html)
            self.assertIn('.workbench-chrome { position: sticky; top: 0;', html)
            self.assertIn("top: calc(var(--sticky-toolbar-height) + var(--sticky-filter-height) - 1px)", html)
            self.assertIn("/api/export_pump_table", html)
            self.assertIn("function exportPumpTable()", html)
            self.assertIn("Array.from(activeFilters.values())", html)

    def test_render_index_uses_generic_title_clickable_project_path_and_hides_manual_refs(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C1,0603,10uF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "project_name": "PXIE_4137 器件选型搜索助手",
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            from html.parser import HTMLParser

            class HeaderParser(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.headers = []

                def handle_starttag(self, tag, attrs):
                    if tag == 'header':
                        self.headers.append(dict(attrs))

            parser = HeaderParser()
            parser.feed(html)
            self.assertEqual(parser.headers, [{'class': 'workbench-chrome'}])
            self.assertIn('class="project-path-link"', html)
            self.assertIn('data-action="open-project-path"', html)
            self.assertLess(html.index('class="app-menubar"'), html.index('class="project-path-link"'))
            self.assertIn("min-height: 27px; padding: 0 8px", html)
            self.assertIn("min-height: 58px; padding: 4px 8px", html)
            self.assertIn("min-height: 31px", html)
            self.assertIn('type="hidden" id="refsInput"', html)

    def test_render_index_has_menus_schematic_refresh_and_default_view_config(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps({"source_bom": "BOM/source.csv", "schematic_dir": "DOC"}),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn('class="app-menubar"', html)
            self.assertIn('id="refreshSchematicRows"', html)
            self.assertIn("刷新原理图当前数据", html)
            self.assertIn('id="saveDefaultView"', html)
            self.assertIn('id="loadDefaultView"', html)
            self.assertIn("DEFAULT_VIEW_STORAGE_KEY", html)
            self.assertIn('let selectedBom = data.default_bom', html)
            self.assertIn('sessionStorage.removeItem("partSelectionSourceOnce")', html)
            self.assertIn('bomSelect.value = "__schematic__"', html)
            self.assertIn("function saveDefaultViewConfig()", html)
            self.assertIn("function loadDefaultViewConfig()", html)
            self.assertNotIn("<textarea id=\"refsInput\"", html)
            self.assertNotIn('id="reloadRows"', html)
            self.assertNotIn("覆盖记录：", html)
            self.assertIn("/api/open_project_path", html)

    def test_render_index_allows_candidate_columns_to_shrink_and_wrap(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Resistor,Resistor,R1,0402,100k\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn("min-width: 28px", html)
            self.assertIn("overflow-wrap: anywhere", html)
            self.assertIn("Math.max(28", html)
            self.assertNotIn('data-candidate-col="param"] { min-width: 520px; max-width', html)

    def test_render_index_has_schematic_sync_button_and_field_diff_rendering(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            (bom_dir / "source.csv").write_text(
                "Comment,Description,Designator,Footprint,Value\n"
                "Capacitor,Capacitor,C118,0402,100pF\n",
                encoding="utf-8",
            )
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            from part_selection_server import render_index

            html = render_index(PartSelectionApp(config_path))

            self.assertIn('data-action="sync-schematic"', html)
            self.assertIn('id="syncAllSchematic"', html)
            self.assertIn("function syncAllSchematic", html)
            self.assertIn("const lcsc_c_code = currentLcscCode(tr)", html)
            self.assertIn("field_overrides.lcsc_c_code = lcsc_c_code", html)
            self.assertIn("field_overrides.lcsc_url = lcscUrl(lcsc_c_code)", html)
            self.assertIn("const description = schematicDescription(tr)", html)
            self.assertIn("field_overrides.description = description", html)
            self.assertNotIn('id="saveRefs"', html)
            self.assertIn("function fieldDiff", html)
            self.assertIn('postSchematicSync("/api/bulk_sync_schematic"', html)

    def test_kicad_search_script_defaults_to_fast_search_panel(self):
        script = (TOOLS_DIR / "kicad_schematic_search.ps1").read_text(encoding="utf-8")

        self.assertIn("[switch]$ReloadFromDisk", script)
        self.assertIn("[switch]$NotifyOnly", script)
        self.assertIn('[string]$KicadSchematicExe = "C:\\Program Files\\KiCad\\10.0\\bin\\eeschema.exe"', script)
        self.assertIn("[switch]$EnablePageSwitch", script)
        self.assertIn('[string]$Mode = "SearchPanel"', script)
        self.assertIn('$EnablePageSwitch -and $Mode -eq "SearchPanel" -and $page', script)
        self.assertIn("Start-Sleep -Milliseconds 420", script)
        self.assertIn("Resolve-SearchPanelScreenPoint", script)
        self.assertIn("SEARCH_PANEL_GRID_RECT source=UIA", script)
        self.assertIn("GetDpiForWindow", script)
        self.assertIn("SEARCH_PANEL_COORDS_DPI", script)
        self.assertIn("SEARCH_PANEL_COORDS_ADJUST", script)
        self.assertIn("SEARCH_PANEL_ACTIVATE method=first_click", script)
        self.assertIn("SEARCH_PANEL_ACTIVATE method=second_click", script)
        self.assertIn('[System.Windows.Forms.SendKeys]::SendWait("{ENTER}")', script)
        self.assertIn("Invoke-ProtectedSearchResultClicks -X $firstResultX -Y $firstResultY -IntervalMs 460", script)
        self.assertNotIn("Wait-FirstSearchPanelResultPoint", script)
        self.assertNotIn("SendMessageRect", script)
        self.assertIn("function Get-KiCadProjectFile", script)
        self.assertIn("function Get-TopLevelSchematicFile", script)
        self.assertIn("function Wait-KiCadEditorFrameClosed", script)
        self.assertIn("CLOSE_MAIN_WINDOW requested=", script)
        self.assertIn("KiCad schematic window did not close after save", script)
        self.assertNotIn("DONE action=save_only no_schematic_window", script)
        save_window_recovery = script[script.index('if ($page -and $windows.Count -eq 0 -and -not $DryRun)'):]
        self.assertIn('Open-SchematicPage -Page $page', save_window_recovery)
        self.assertIn('Wait-KiCadWindow -Page "" -TimeoutMs 20000', save_window_recovery)
        no_window_guard = save_window_recovery[save_window_recovery.index('if ($windows.Count -eq 0)'):]
        self.assertIn('External schematic write was aborted', no_window_guard)
        self.assertIn('throw $message', no_window_guard)
        self.assertNotIn('if ($SaveOnly)', no_window_guard[:no_window_guard.index('Ensure-NativeMethods')])
        self.assertIn("OPEN_PROJECT $projectPath", script)
        self.assertIn('Start-Process -FilePath $KicadExe -ArgumentList @(\'"\' + $projectPath + \'"\')', script)
        self.assertNotIn('Start-Process -FilePath $KicadSchematicExe', script)
        self.assertIn('OPEN_PAGE via_manager', script)
        self.assertIn('Find-KiCadManagerSchematicCommand', script)
        self.assertNotIn("Start-Process -FilePath $KicadSchematicExe -ArgumentList @($schematicPath)", script)
        self.assertIn("OPEN_PAGE top_schematic=$topSchematicPath target_page=$Page", script)
        self.assertNotIn('Start-Process -FilePath $KicadExe -ArgumentList "`"$schematicPath`""', script)
        self.assertIn("OPEN_PAGE no_schematic_window", script)
        self.assertNotIn("OPEN_PAGE skipped because KiCad is already running", script)
        self.assertIn('Write-Log "DONE action=notify_only"', script)
        self.assertNotIn('Invoke-KiCadApi -ProcessId $target.Id -ApiName "revert"', script)

    def test_save_group_manual_c_code_applies_to_all_refs(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "source_bom": "BOM/source.csv",
                        "schematic_dir": "DOC",
                        "output_overrides": "BOM/overrides.json",
                    }
                ),
                encoding="utf-8",
            )

            app = PartSelectionApp(config_path)
            app.save_override(
                {
                    "refs": "R1,R2",
                    "accepted_c_code": "C49678",
                    "manual_c_code": "C49678",
                    "decision": "AcceptManual",
                    "status": "ChangedToNewCode",
                }
            )

            overrides = load_overrides(bom_dir / "overrides.json")
            self.assertEqual(overrides["R1"]["manual_c_code"], "C49678")
            self.assertEqual(overrides["R2"]["manual_c_code"], "C49678")
            self.assertEqual(overrides["R1"]["status"], "ChangedToNewCode")


class CandidateSelectionRegressionTest(unittest.TestCase):
    def test_select_candidate_saves_group_in_one_override_write(self):
        with temporary_project() as tmp:
            root = Path(tmp)
            bom_dir = root / "BOM"
            bom_dir.mkdir()
            config_path = bom_dir / "part_selection_config.json"
            config_path.write_text(json.dumps({
                "schematic_dir": "DOC", "output_overrides": "BOM/overrides.json",
                "stock_cache": "BOM/stock.json",
            }), encoding="utf-8")
            app = PartSelectionApp(config_path)
            from part_selection_server import save_overrides as actual_save
            with mock.patch("part_selection_server.save_overrides", wraps=actual_save) as write:
                result = app.select_candidate({
                    "refs": "C1,C2,C3", "candidate": {
                        "c_code": "C88946", "model": "GRM155R71E474KE14D",
                        "package": "0402", "param_detail": "容值: 470nF; 额定电压: 25V",
                    },
                    "replace_value": True,
                })
            self.assertEqual(write.call_count, 1)
            self.assertEqual(result["refs"], ["C1", "C2", "C3"])
            saved = load_overrides(bom_dir / "overrides.json")
            self.assertEqual([saved[ref]["manual_c_code"] for ref in result["refs"]], ["C88946"] * 3)
            self.assertEqual([saved[ref]["field_overrides"]["value"] for ref in result["refs"]], ["470nF"] * 3)

    def test_candidate_value_strips_search_highlight_markup(self):
        fields = part_selection_server.candidate_selection_fields({
            "model": "GRM155R71E474KE14D", "package": "0402", "c_code": "C88946",
            "param_detail": "容值: <SPAN CLASS='LUCENE_HIGHLIGHT_CLASS'>470nF</SPAN>; 额定电压: 25V",
        }, ["C1688"])
        self.assertEqual(fields["value"], "470nF")


if __name__ == "__main__":
    unittest.main()
