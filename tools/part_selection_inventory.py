"""Local inventory import and comparison, kept separate from marketplace stock."""
from __future__ import annotations

import json
import copy
import hashlib
import os
import re
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

from part_selection_overrides import normalize_lcsc_c_code


ALIASES = {
    "c_code": ("商品编号", "商品编码", "C编码", "C 编码", "料号", "LCSC编号"),
    "model": ("厂家型号", "型号", "制造商型号"),
    "brand": ("品牌", "厂商品牌"),
    "package": ("封装", "封装规格"),
    "order": ("订单编号", "订单号"),
    "quantity": ("同型号已购合计", "已购合计", "现存数量"),
}


def _header(value: Any) -> str:
    return re.sub(r"[\s_]", "", str(value or "").lower())


def _column_map(values) -> dict[str, int]:
    headers = {_header(value): index for index, value in enumerate(values) if value is not None}
    result = {}
    for field, aliases in ALIASES.items():
        found = next((headers[_header(alias)] for alias in aliases if _header(alias) in headers), None)
        if found is not None:
            result[field] = found
    return result


def parse_inventory_workbook(path: Path, sheet_name: str = "") -> dict[str, object]:
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("本地库存仅支持 .xlsx 或 .xlsm 文件")
    if not path.is_file():
        raise FileNotFoundError(f"找不到库存文件：{path}")
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("库存文件超过 20 MB，请先拆分")
    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("库存工作簿解压后超过 100 MB，请先拆分")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        if sheet_name:
            if sheet_name not in names:
                raise ValueError("指定的工作表不存在")
            sheet = workbook[sheet_name]
        else:
            sheet = next((item for item in workbook.worksheets
                          if "c_code" in _column_map(next(item.iter_rows(min_row=1, max_row=1, values_only=True), ()))
                          and "quantity" in _column_map(next(item.iter_rows(min_row=1, max_row=1, values_only=True), ()))), None)
            if sheet is None:
                return {"sheets": names, "sheet": "", "items": {}, "unmatched": {}, "issues": ["没有找到同时包含商品编号和同型号已购合计的工作表"],
                        "summary": {"ready": 0, "unresolved": 0, "unmatched": 0, "invalid": 1}}
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = _column_map(headers)
        if "c_code" not in columns or "quantity" not in columns:
            raise ValueError("所选工作表必须包含“商品编号”和“同型号已购合计”列")

        grouped: dict[str, list[dict[str, Any]]] = {}
        unmatched_groups: dict[str, list[dict[str, Any]]] = {}
        issues: list[str] = []
        invalid_count = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_number > 50001:
                raise ValueError("库存明细超过 50,000 行，请先拆分")
            raw = {key: (values[index] if index < len(values) else None) for key, index in columns.items()}
            if not any(value is not None and str(value).strip() for value in raw.values()):
                continue
            raw_code = str(raw.get("c_code") or "").strip()
            code = normalize_lcsc_c_code(raw_code) if re.fullmatch(r"C\s*\d+", raw_code, re.I) else ""
            model = str(raw.get("model") or "").strip()
            entry = {
                "row": row_number, "model": model,
                "brand": str(raw.get("brand") or "").strip(),
                "package": str(raw.get("package") or "").strip(),
                "order": str(raw.get("order") or "").strip(), "source_product_code": raw_code,
            }
            unmatched_key = ''
            if not code and model and (not raw_code or re.fullmatch(r"[A-BD-Z]\d+", raw_code, re.I)):
                unmatched_key = hashlib.sha256(json.dumps([entry[field].upper() for field in ('model', 'brand', 'package')]).encode()).hexdigest()[:24]
            try:
                quantity = int(raw.get("quantity"))
                if quantity < 0 or str(raw.get("quantity")).strip() != str(quantity):
                    raise ValueError
            except (TypeError, ValueError, OverflowError):
                invalid_count += 1
                issues.append(f"第 {row_number} 行：同型号已购合计不是有效的非负整数")
                entry.update({'invalid': True, 'quantity': None})
                if code:
                    grouped.setdefault(code, []).append(entry)
                elif unmatched_key:
                    unmatched_groups.setdefault(unmatched_key, []).append(entry)
                continue
            entry['quantity'] = quantity
            if unmatched_key:
                unmatched_groups.setdefault(unmatched_key, []).append(entry)
                continue
            if not code:
                invalid_count += 1
                issues.append(f"第 {row_number} 行：商品编号不是有效的 C 编码")
                continue
            grouped.setdefault(code, []).append(entry)

        items = {}
        for code, entries in grouped.items():
            if any(entry.get("invalid") for entry in entries):
                items[code] = {"c_code": code, "quantity": None, "status": "unresolved",
                               "issue": "含无效汇总数量的来源行", "source_rows": [entry["row"] for entry in entries]}
                continue
            quantities = {entry["quantity"] for entry in entries}
            base = entries[0]
            items[code] = {
                "c_code": code,
                "quantity": next(iter(quantities)) if len(quantities) == 1 else None,
                "status": "ready" if len(quantities) == 1 else "unresolved",
                "issue": "" if len(quantities) == 1 else "同一 C 编码的同型号已购合计不一致",
                "model": base["model"], "brand": base["brand"], "package": base["package"],
                "source_orders": sorted({entry["order"] for entry in entries if entry["order"]}),
                "source_rows": [entry["row"] for entry in entries],
            }
            if len(quantities) != 1:
                issues.append(f"{code}：{items[code]['issue']}（第 {', '.join(map(str, items[code]['source_rows']))} 行）")
        unmatched = {}
        for key, entries in unmatched_groups.items():
            quantities = {entry['quantity'] for entry in entries}
            valid = len(quantities) == 1 and None not in quantities
            base = entries[0]
            unmatched[key] = {field: base[field] for field in ('model', 'brand', 'package')}
            unmatched[key].update({'quantity': next(iter(quantities)) if valid else None,
                'status': 'unmatched' if valid else 'unresolved',
                'issue': '需要明确指定对应主料 C 编码' if valid else '同型号汇总量冲突或无效，需核实',
                'source_product_codes': sorted({entry['source_product_code'] for entry in entries if entry['source_product_code']}),
                'source_orders': sorted({entry['order'] for entry in entries if entry['order']}),
                'source_rows': [entry['row'] for entry in entries]})
        summary = {
            "ready": sum(item["status"] == "ready" for item in items.values()),
            "unresolved": sum(item["status"] == "unresolved" for item in items.values()),
            "invalid": invalid_count,
            "unmatched": len(unmatched),
        }
        return {"sheets": names, "sheet": sheet.title, "items": items, "unmatched": unmatched, "issues": issues, "summary": summary,
                "source_name": path.name}
    finally:
        workbook.close()


def merge_inventory_records(existing: dict, preview: dict) -> dict:
    result = {"schema_version": 1, "items": dict(existing.get("items") or {})}
    owners = inventory_source_owners(existing)
    for code, item in preview.get('items', {}).items():
        for key in item.get('source_group_keys', []):
            if key in owners and owners[key] != code:
                raise ValueError('同一库存来源已经绑定其他 C 编码，不能重复归属')
    unmatched = {**existing.get('unmatched', {}), **preview.get('unmatched', {})}
    for key in preview.get('matched_keys', []):
        unmatched.pop(key, None)
    if unmatched:
        result['unmatched'] = unmatched
    for code, incoming in (preview.get("items") or {}).items():
        previous = result["items"].get(code, {})
        result["items"][code] = {**previous, **incoming}
        keys = sorted(set(previous.get('source_group_keys', []) + incoming.get('source_group_keys', [])))
        if keys:
            result['items'][code]['source_group_keys'] = keys
    return result


def inventory_source_owners(inventory: dict) -> dict:
    owners = {}
    for code, item in inventory.get('items', {}).items():
        for key in item.get('source_group_keys', []):
            if key in owners and owners[key] != code:
                raise ValueError('同一库存来源重复绑定，请先核对库存文件')
            owners[key] = code
    return owners


def apply_inventory_matches(preview: dict, matches: dict) -> dict:
    result = copy.deepcopy(preview)
    if not isinstance(matches, dict) or set(matches) - result.get('unmatched', {}).keys():
        raise ValueError('库存匹配包含未知记录，请重新预览')
    result['match_candidates'] = copy.deepcopy(result.get('unmatched', {}))
    result['matches'] = dict(matches)
    result['matched_keys'] = []
    for key, code in matches.items():
        if not isinstance(code, str) or not re.fullmatch(r'C\d+', code, re.I):
            raise ValueError('请为待匹配库存填写有效主料 C 编码')
        code = code.upper()
        incoming = result['unmatched'].pop(key)
        incoming['source_group_keys'] = [key]
        incoming.update({'c_code': code, 'status': 'ready' if incoming['quantity'] is not None else 'unresolved',
                         'issue': '' if incoming['quantity'] is not None else incoming['issue']})
        old = result['items'].get(code)
        if old:
            if old.get('status') != 'ready' or incoming['status'] != 'ready' or old['quantity'] != incoming['quantity']:
                incoming.update({'quantity': None, 'status': 'unresolved', 'issue': '手动匹配与同编码汇总数量冲突，未计入可用库存'})
                result['issues'].append(f'{code}：{incoming["issue"]}')
            for field in ('source_orders', 'source_rows', 'source_product_codes', 'source_group_keys'):
                incoming[field] = sorted(set(old.get(field, []) + incoming.get(field, [])))
        result['items'][code] = incoming
        result['matched_keys'].append(key)
    result['summary'].update({'ready': sum(item['status'] == 'ready' for item in result['items'].values()),
        'unresolved': sum(item['status'] == 'unresolved' for item in result['items'].values()),
        'unmatched': len(result.get('unmatched', {}))})
    return result


def load_inventory(path: Path) -> dict:
    path = Path(path)
    if not path.exists():
        return {"schema_version": 1, "items": {}}
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict) or data.get("schema_version") != 1 or not isinstance(data.get("items"), dict):
        raise ValueError("本地库存文件格式无效")
    for code, record in data["items"].items():
        quantity = record.get("quantity") if isinstance(record, dict) else None
        if (not re.fullmatch(r"C\d+", code) or not isinstance(record, dict)
                or (record.get("status", "ready") == "ready"
                    and (isinstance(quantity, bool) or not isinstance(quantity, int) or quantity < 0))):
            raise ValueError(f"本地库存记录无效：{code}")
    if 'unmatched' in data and (not isinstance(data['unmatched'], dict) or any(not isinstance(item, dict) for item in data['unmatched'].values())):
        raise ValueError('待匹配本地库存格式无效')
    return {"schema_version": 1, "items": data['items'], **({'unmatched': data['unmatched']} if data.get('unmatched') else {})}


def save_inventory(path: Path, inventory: dict) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent,
                                         prefix=".inventory-", suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
            json.dump(inventory, stream, ensure_ascii=False, indent=2, sort_keys=True)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if temporary and temporary.exists():
            temporary.unlink()


def calculate_inventory_status(record: dict | None, per_set_quantity: int, set_count: int) -> dict:
    if isinstance(set_count, bool) or not isinstance(set_count, int) or not 1 <= set_count <= 1_000_000:
        raise ValueError("套数必须为 1 到 1,000,000 的正整数")
    if isinstance(per_set_quantity, bool) or not isinstance(per_set_quantity, int) or per_set_quantity < 0:
        raise ValueError("单套用量必须为非负整数")
    if not record:
        return {"inventory": None, "per_set_quantity": per_set_quantity, "set_count": set_count,
                "required": per_set_quantity * set_count, "difference": None, "status": "unmatched"}
    if (record.get("status", "ready") != "ready" or isinstance(record.get("quantity"), bool)
            or not isinstance(record.get("quantity"), int) or record["quantity"] < 0):
        return {"inventory": None, "per_set_quantity": per_set_quantity, "set_count": set_count,
                "required": per_set_quantity * set_count, "difference": None, "status": "unresolved"}
    required = per_set_quantity * set_count
    quantity = record["quantity"]
    return {"inventory": quantity, "per_set_quantity": per_set_quantity, "set_count": set_count,
            "required": required, "difference": quantity - required,
            "status": "sufficient" if quantity >= required else "insufficient"}


def inventory_import_time() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")
